from __future__ import annotations

import json
import os
import subprocess
import sys
import tempfile
import textwrap
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]

Report = dict[str, list[tuple[object, ...]]]


def _environment(**updates: str) -> dict[str, str]:
    environment = {
        key: value
        for key, value in os.environ.items()
        if not key.startswith("ATEN_SIZE_PROFILER_")
    }
    environment["PYTHONPATH"] = str(ROOT)
    environment.update(updates)
    return environment


def _run(name: str, source: str, **updates: str) -> None:
    result = subprocess.run(
        [sys.executable, "-c", textwrap.dedent(source)],
        cwd=ROOT,
        env=_environment(**updates),
        text=True,
        capture_output=True,
        timeout=240,
    )
    assert result.returncode == 0, (
        f"{name} failed with code {result.returncode}\n"
        f"stdout:\n{result.stdout}\nstderr:\n{result.stderr}"
    )


def _read_report(path: Path) -> Report:
    workbook = load_workbook(path, read_only=True)
    try:
        return {
            name: list(workbook[name].iter_rows(values_only=True))
            for name in workbook.sheetnames
        }
    finally:
        workbook.close()


def _sheet_rows(report: Report, sheet: str) -> list[dict[str, object]]:
    header, *rows = report[sheet]
    return [dict(zip(header, row)) for row in rows]


def _metadata(report: Report) -> dict[str, object]:
    return {row[0]: row[1] for row in report["metadata"][1:]}


# P1-1: a failed optimizer.step() in the target iteration must not leave its
# data behind when it is retried. This must hold whether the profiler is
# imported before or after the optimizer and its LR scheduler (a scheduler
# replaces step() on the instance), after stop/start, and for an optimizer
# that reuses the id of a freed one.
#
#   FAIL=none     batches 2, 4, 5: the baseline.
#   FAIL=forward  batches 2, 3, 4, 5: step() of batch 3 raises before
#                 updating and the loop starts over with batch 4.
#   FAIL=step     batches 2, 4, 5: step() of batch 4 raises after updating
#                 and only step() is retried.
#
# Iteration 1 is batch 4 in every case, so the reports must be identical.
TRAINING_SOURCE = """
import gc
import os
import warnings
import weakref

import torch
from torch import nn

ORDER = os.environ["ORDER"]
FAIL = os.environ["FAIL"]


class FlakySGD(torch.optim.SGD):
    calls = 0

    def step(self, closure=None):
        FlakySGD.calls += 1
        fail = FlakySGD.calls == 2
        if fail and FAIL == "forward":
            raise RuntimeError("transient optimizer failure")
        result = super().step(closure)
        if fail and FAIL == "step":
            raise RuntimeError("transient optimizer failure")
        return result


RAW_STEP = FlakySGD.__dict__["step"]
model = nn.Linear(4, 4)


def new_optimizer():
    return FlakySGD(model.parameters(), lr=0.1)


def new_scheduler(optimizer):
    scheduler = torch.optim.lr_scheduler.StepLR(
        optimizer, step_size=1, gamma=0.5,
    )
    return scheduler, optimizer.__dict__["step"]


def reused_optimizer():
    # CPython gives a freed optimizer's memory to the next one, so the new
    # optimizer has the id of one the profiler has already wrapped.
    for _ in range(100):
        stale = new_optimizer()
        stale_id = id(stale)
        del stale
        optimizer = new_optimizer()
        if id(optimizer) == stale_id:
            return optimizer
        del optimizer
    raise AssertionError("no optimizer reused the id of a freed one")


if ORDER in ("optimizer_first", "scheduler_first"):
    optimizer, spare = new_optimizer(), new_optimizer()
if ORDER == "scheduler_first":
    scheduler, scheduler_step = new_scheduler(optimizer)
    spare_scheduler, _ = new_scheduler(spare)

import aten_size_profiler as profiler

if ORDER in ("profiler_first", "restart"):
    optimizer, spare = new_optimizer(), new_optimizer()
if ORDER == "reuse":
    optimizer, spare = reused_optimizer(), new_optimizer()
if ORDER != "scheduler_first":
    scheduler, scheduler_step = new_scheduler(optimizer)
    spare_scheduler, _ = new_scheduler(spare)
if ORDER == "restart":
    for _ in range(2):
        profiler.stop_profiler(export=False)
        profiler.start_profiler()

# The installed profiler must not keep an optimizer alive.
spare_ref = weakref.ref(spare)
del spare, spare_scheduler
gc.collect()
assert spare_ref() is None

batches = iter((2, 3, 4, 5) if FAIL == "forward" else (2, 4, 5))
completed = 0
with warnings.catch_warnings(record=True) as caught:
    warnings.simplefilter("always")
    while completed < 3:
        optimizer.zero_grad(set_to_none=True)
        model(torch.randn(next(batches), 4)).square().mean().backward()
        try:
            optimizer.step()
        except RuntimeError:
            if FAIL == "forward":
                continue
            optimizer.step()
        scheduler.step()
        completed += 1

assert FlakySGD.calls == (3 if FAIL == "none" else 4), FlakySGD.calls
# The scheduler still sees its own step() wrapper and the optimizer calls.
scheduler_warnings = [
    str(item.message)
    for item in caught
    if "optimizer.step()" in str(item.message)
]
assert not scheduler_warnings, scheduler_warnings
assert scheduler.get_last_lr() == [0.1 * 0.5 ** 3], scheduler.get_last_lr()
assert profiler.get_iteration_status()["phase"] == "FINISHED"
assert profiler.stop_profiler() is not None

# Stop restores the steps the profiler found; PyTorch's own profiling
# wrapper around the class step stays.
assert FlakySGD.__dict__["step"].__wrapped__ is RAW_STEP
assert optimizer.__dict__["step"] is scheduler_step
optimizer_ref = weakref.ref(optimizer)
del optimizer, scheduler, scheduler_step
gc.collect()
assert optimizer_ref() is None
"""

ORDERS = (
    "profiler_first", "optimizer_first", "scheduler_first", "restart", "reuse",
)

# Metadata that differs between runs by construction.
RUN_SPECIFIC_METADATA = {"created_at", "output_file", "configured_output_filename"}


def _run_training(output_dir: Path, order: str, fail: str) -> Report:
    report = output_dir / f"train_{order}_{fail}.xlsx"
    _run(
        f"training {order}/{fail}",
        TRAINING_SOURCE,
        ORDER=order,
        FAIL=fail,
        ATEN_SIZE_PROFILER_ENABLED="true",
        ATEN_SIZE_PROFILER_ITERATION_MODE="optimizer_step",
        ATEN_SIZE_PROFILER_TARGET_ITER="1",
        ATEN_SIZE_PROFILER_OUTPUT_DIR=str(output_dir),
        ATEN_SIZE_PROFILER_OUTPUT_FILENAME=report.name,
    )
    return _read_report(report)


def _check_baseline(report: Report) -> None:
    detail = _sheet_rows(report, "detail")
    assert [row["call_id"] for row in detail] == list(range(len(detail)))
    assert {row["iteration"] for row in detail} == {1}
    input_tensors = _sheet_rows(report, "input_tensors")
    assert len(input_tensors) == sum(row["input_count"] for row in detail)
    assert {row["iteration"] for row in input_tensors} == {1}

    # Only the forward of batch 4; args[1] of addmm is the batch input.
    addmm = [row for row in detail if row["op_name"] == "aten.addmm.default"]
    assert len(addmm) == 1, addmm
    assert json.loads(addmm[0]["input_shapes"])[1] == [4, 4], addmm[0]
    # One SGD update per parameter, at the lr after one scheduler step.
    updates = [row for row in detail if row["op_name"] == "aten.add_.Tensor"]
    assert len(updates) == 2, updates
    for row in updates:
        assert json.loads(row["scalar_args"]) == [["kwargs['alpha']", "-0.05"]]

    detail_counts = Counter(row["op_name"] for row in detail)
    unique_shapes = _sheet_rows(report, "unique_shapes")
    shape_counts: Counter[object] = Counter()
    for row in unique_shapes:
        shape_counts[row["op_name"]] += row["count"]
    summary = _sheet_rows(report, "summary")
    assert {row["op_name"] for row in summary} == set(detail_counts)
    for row in summary:
        op_name = row["op_name"]
        assert row["call_count"] == detail_counts[op_name], row
        assert row["call_count"] == shape_counts[op_name], row

    metadata = _metadata(report)
    assert metadata["call_count"] == len(detail), metadata
    assert metadata["iteration_complete"] is True, metadata
    assert metadata["iteration_phase"] == "FINISHED", metadata
    assert metadata["iteration_mode"] == "optimizer_step", metadata
    assert metadata["target_iter"] == 1, metadata


def _assert_same_report(name: str, actual: Report, expected: Report) -> None:
    assert list(actual) == list(expected), (name, list(actual))
    for sheet in expected:
        if sheet == "metadata":
            actual_rows: object = {
                key: value
                for key, value in _metadata(actual).items()
                if key not in RUN_SPECIFIC_METADATA
            }
            expected_rows: object = {
                key: value
                for key, value in _metadata(expected).items()
                if key not in RUN_SPECIFIC_METADATA
            }
        else:
            actual_rows = actual[sheet]
            expected_rows = expected[sheet]
        assert actual_rows == expected_rows, (
            f"{name}: sheet {sheet!r} differs from the baseline\n"
            f"actual:   {actual_rows}\nexpected: {expected_rows}"
        )


def _check_optimizer_retry(output_dir: Path) -> None:
    cases = [
        (order, fail)
        for order in ORDERS
        for fail in ("none", "forward", "step")
    ]
    # Each case is a separate interpreter; run them side by side.
    with ThreadPoolExecutor(max_workers=4) as executor:
        reports = dict(zip(
            cases,
            executor.map(lambda case: _run_training(output_dir, *case), cases),
        ))
    baseline = reports["profiler_first", "none"]
    _check_baseline(baseline)
    for (order, fail), report in reports.items():
        _assert_same_report(f"{order}/{fail}", report, baseline)


# P1-2: mp.spawn workers without RANK and without any ATen operator between
# init_process_group and destroy_process_group.
SPAWN_SOURCE = """
import os
import sys

import torch.distributed as dist
import torch.multiprocessing as mp


def worker(rank, world_size, init_file):
    import aten_size_profiler

    dist.init_process_group(
        "gloo",
        init_method=f"file://{init_file}",
        rank=rank,
        world_size=world_size,
    )
    dist.destroy_process_group()


if __name__ == "__main__":
    for key in ("RANK", "WORLD_SIZE", "LOCAL_RANK"):
        os.environ.pop(key, None)
    mp.spawn(worker, args=(2, sys.argv[1]), nprocs=2, join=True)
"""


def _check_spawn_without_operators(output_dir: Path) -> None:
    init_file = output_dir / "spawn_init"
    report_dir = output_dir / "spawn"
    script = output_dir / "spawn_worker.py"
    script.write_text(textwrap.dedent(SPAWN_SOURCE), encoding="utf-8")
    result = subprocess.run(
        [sys.executable, str(script), str(init_file)],
        cwd=ROOT,
        env=_environment(
            ATEN_SIZE_PROFILER_ENABLED="true",
            ATEN_SIZE_PROFILER_OUTPUT_DIR=str(report_dir),
            ATEN_SIZE_PROFILER_OUTPUT_FILENAME="empty.xlsx",
        ),
        text=True,
        capture_output=True,
        timeout=240,
    )
    assert result.returncode == 0, (
        f"mp.spawn failed with code {result.returncode}\n"
        f"stdout:\n{result.stdout}\nstderr:\n{result.stderr}"
    )
    files = sorted(path.name for path in report_dir.iterdir())
    assert files == ["empty_rank0.xlsx", "empty_rank1.xlsx"], files
    for rank in (0, 1):
        metadata = _metadata(
            _read_report(report_dir / f"empty_rank{rank}.xlsx")
        )
        assert metadata["global_rank"] == rank, metadata
        assert metadata["world_size"] == 2, metadata
        assert metadata["distributed_initialized"] is True, metadata
        assert metadata["call_count"] == 0, metadata


def run() -> None:
    with tempfile.TemporaryDirectory(prefix="aten-codex-p1-") as path:
        output_dir = Path(path)
        _check_optimizer_retry(output_dir)
        _check_spawn_without_operators(output_dir)
    print("CODEX P1 CHECK PASS")


if __name__ == "__main__":
    run()
