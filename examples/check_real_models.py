from __future__ import annotations

import argparse
import json
import os
import subprocess
import sys
import time
from collections import Counter
from pathlib import Path
from statistics import median
from typing import Any

import torch
from openpyxl import load_workbook
from torch import nn


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = (
    ROOT / "aten_size_reports" / "stage7_validation" / f"run_{os.getpid()}"
)
EXPECTED_SHEETS = [
    "summary",
    "unique_shapes",
    "detail",
    "input_tensors",
    "metadata",
]


class TinyVisionNet(nn.Module):
    def __init__(self) -> None:
        super().__init__()
        self.features = nn.Sequential(
            nn.Conv2d(3, 8, kernel_size=3, padding=1),
            nn.BatchNorm2d(8),
            nn.ReLU(),
            nn.Conv2d(8, 16, kernel_size=3, padding=1),
            nn.ReLU(),
            nn.AdaptiveAvgPool2d(1),
        )
        self.classifier = nn.Linear(16, 4)

    def forward(self, inputs: torch.Tensor) -> torch.Tensor:
        features = self.features(inputs)
        return self.classifier(features.flatten(1))


class TinyTransformer(nn.Module):
    def __init__(self) -> None:
        super().__init__()
        layer = nn.TransformerEncoderLayer(
            d_model=16,
            nhead=4,
            dim_feedforward=32,
            batch_first=True,
        )
        self.encoder = nn.TransformerEncoder(layer, num_layers=2)
        self.classifier = nn.Linear(16, 4)

    def forward(self, inputs: torch.Tensor) -> torch.Tensor:
        encoded = self.encoder(inputs)
        return self.classifier(encoded.mean(dim=1))


def _build_model(model_name: str, device: torch.device) -> nn.Module:
    if model_name == "vision":
        model = TinyVisionNet()
    elif model_name == "transformer":
        model = TinyTransformer()
    else:
        raise ValueError(f"unknown model: {model_name}")
    return model.to(device)


def _make_input(
    model_name: str,
    batch_size: int,
    device: torch.device,
) -> torch.Tensor:
    if model_name == "vision":
        shape = (batch_size, 3, 8, 8)
    else:
        shape = (batch_size, 5, 16)
    return torch.randn(*shape, device=device)


def _model_input_shape(model_name: str) -> str:
    if model_name == "vision":
        return "[3,3,8,8]"
    return "[3,5,16]"


def _expected_operators(model_name: str) -> set[str]:
    if model_name == "vision":
        return {"aten.convolution.default", "aten.addmm.default"}
    return {"aten.mm.default", "aten.addmm.default"}


def _collector_bytes(collector: Any) -> int:
    total = sys.getsizeof(collector.calls) + sys.getsizeof(
        collector.aggregates
    )
    for call in collector.calls:
        total += sys.getsizeof(call)
        total += sys.getsizeof(call.op_name)
        total += sys.getsizeof(call.inputs)
        total += sys.getsizeof(call.stack)
        for tensor in call.inputs:
            total += sys.getsizeof(tensor)
            total += sys.getsizeof(tensor.path)
            total += sys.getsizeof(tensor.shape)
            total += sys.getsizeof(tensor.dtype)
            total += sys.getsizeof(tensor.device)
            total += sys.getsizeof(tensor.stride)
    for aggregate in collector.aggregates.values():
        total += sys.getsizeof(aggregate)
        total += sys.getsizeof(aggregate.op_name)
        total += sys.getsizeof(aggregate.input_signature_counts)
    return total


def _run_workload(
    model_name: str,
    mode_name: str,
    device_name: str,
    *,
    profiled: bool,
    benchmark: bool,
    output_path: Path | None = None,
) -> dict[str, Any]:
    environment = {
        key: value
        for key, value in os.environ.items()
        if not key.startswith("ATEN_SIZE_PROFILER_")
    }
    environment.update(
        {
            "PYTHONPATH": str(ROOT),
            "TEST_MODEL": model_name,
            "TEST_MODE": mode_name,
            "TEST_DEVICE": device_name,
            "PROFILED": str(profiled).lower(),
            "BENCHMARK": str(benchmark).lower(),
        }
    )
    if profiled:
        environment.update(
            {
                "ATEN_SIZE_PROFILER_ENABLED": "true",
                "ATEN_SIZE_PROFILER_OUTPUT_DIR": str(REPORT_DIR),
                "ATEN_SIZE_PROFILER_OUTPUT_FILENAME": (
                    output_path.name if output_path is not None else ""
                ),
            }
        )
        if not benchmark:
            environment.update(
                {
                    "ATEN_SIZE_PROFILER_TARGET_ITER": "1",
                    "ATEN_SIZE_PROFILER_ITERATION_MODE": (
                        "forward"
                        if mode_name == "inference"
                        else "optimizer_step"
                    ),
                }
            )
    result = subprocess.run(
        [sys.executable, __file__, "--child"],
        cwd=ROOT,
        env=environment,
        text=True,
        capture_output=True,
    )
    assert result.returncode == 0, (
        f"{model_name}-{mode_name}-{device_name} child failed with "
        f"code {result.returncode}\nstdout:\n{result.stdout}\n"
        f"stderr:\n{result.stderr}"
    )
    marker = "BENCHMARK=" if benchmark else "RESULTS="
    lines = [
        line.removeprefix(marker)
        for line in result.stdout.splitlines()
        if line.startswith(marker)
    ]
    assert len(lines) == 1, (
        f"missing {marker} from {model_name}-{mode_name}-{device_name}: "
        f"{result.stdout}\n{result.stderr}"
    )
    return json.loads(lines[0])


def _validate_report(
    path: Path,
    model_name: str,
    mode_name: str,
) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    assert workbook.sheetnames == EXPECTED_SHEETS
    rows = {
        sheet_name: list(
            workbook[sheet_name].iter_rows(min_row=2, values_only=True)
        )
        for sheet_name in EXPECTED_SHEETS
    }
    summary_counts = {row[0]: row[1] for row in rows["summary"]}
    detail_counts = Counter(row[1] for row in rows["detail"])
    unique_counts: Counter[str] = Counter()
    for row in rows["unique_shapes"]:
        unique_counts[row[0]] += row[5]

    assert rows["detail"]
    assert rows["input_tensors"]
    assert detail_counts == summary_counts
    assert unique_counts == summary_counts
    assert sum(row[1] for row in rows["summary"]) == len(rows["detail"])
    assert sum(row[5] for row in rows["unique_shapes"]) == len(
        rows["detail"]
    )
    assert sum(row[3] for row in rows["summary"]) == len(
        rows["input_tensors"]
    )
    assert [row[0] for row in rows["detail"]] == list(
        range(len(rows["detail"]))
    )
    assert {row[12] for row in rows["detail"]} == {1}
    assert {row[11] for row in rows["input_tensors"]} == {1}

    input_shapes = {row[4] for row in rows["input_tensors"]}
    assert _model_input_shape(model_name) in input_shapes
    operators = set(detail_counts)
    assert operators.intersection(_expected_operators(model_name))
    if mode_name == "training":
        assert any("backward" in name for name in operators)
        assert operators.intersection(
            {"aten.add_.Tensor", "aten._foreach_add_.List"}
        )

    metadata = dict(rows["metadata"])
    assert metadata["call_count"] == len(rows["detail"])
    assert metadata["target_iter"] == 1
    assert metadata["iteration_mode"] == (
        "forward" if mode_name == "inference" else "optimizer_step"
    )
    workbook.close()


def _child_run() -> None:
    model_name = os.environ["TEST_MODEL"]
    mode_name = os.environ["TEST_MODE"]
    device = torch.device(os.environ["TEST_DEVICE"])
    profiled = os.environ["PROFILED"] == "true"
    benchmark = os.environ["BENCHMARK"] == "true"

    torch.manual_seed(2027)
    model = _build_model(model_name, device)
    optimizer = None
    if mode_name == "training":
        model.train()
        optimizer = torch.optim.SGD(model.parameters(), lr=0.01)
    else:
        model.eval()

    profiler = None
    if profiled:
        import aten_size_profiler as profiler_module

        profiler = profiler_module

    batch_sizes = [2, 3, 4]
    inputs = [
        _make_input(model_name, batch_size, device)
        for batch_size in batch_sizes
    ]
    results: list[list[Any]] = []

    def step(input_tensor: torch.Tensor) -> None:
        if mode_name == "inference":
            with torch.no_grad():
                output = model(input_tensor)
        else:
            assert optimizer is not None
            optimizer.zero_grad(set_to_none=True)
            output = model(input_tensor)
            output.square().mean().backward()
            optimizer.step()
        results.append(output.detach().cpu().tolist())

    if benchmark:
        warmup_steps = 3
        rounds = 5
        steps_per_round = 8
        for _ in range(warmup_steps):
            step(inputs[0])
        durations = []
        for _ in range(rounds):
            if device.type == "cuda":
                torch.cuda.synchronize()
            start_time = time.perf_counter()
            for _ in range(steps_per_round):
                step(inputs[0])
            if device.type == "cuda":
                torch.cuda.synchronize()
            durations.append(time.perf_counter() - start_time)
        payload = {
            "elapsed_seconds": median(durations),
            "durations": durations,
            "measurement_mode": "continuous_collection",
            "warmup_steps": warmup_steps,
            "rounds": rounds,
            "steps_per_round": steps_per_round,
            "calls": len(profiler.get_collector().calls)
            if profiler is not None
            else 0,
            "collector_python_bytes": _collector_bytes(
                profiler.get_collector()
            )
            if profiler is not None
            else 0,
        }
        if profiler is not None:
            profiler.stop_profiler(export=False)
        print(f"BENCHMARK={json.dumps(payload)}")
        return

    for input_tensor in inputs:
        step(input_tensor)
    parameters = (
        [
            parameter.detach().cpu().tolist()
            for parameter in model.parameters()
        ]
        if mode_name == "training"
        else []
    )
    if not profiled:
        print(
            "RESULTS="
            + json.dumps({"outputs": results, "parameters": parameters})
        )
        return
    assert profiler is not None
    assert profiler.get_iteration_status()["phase"] == "FINISHED"
    report = profiler.stop_profiler()
    assert report is not None
    print(
        "RESULTS="
        + json.dumps(
            {
                "outputs": results,
                "parameters": parameters,
                "report": str(report),
            }
        )
    )


def _run_validation() -> None:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    devices = ["cpu"] + (["cuda"] if torch.cuda.is_available() else [])
    validation_results: list[str] = []
    benchmark_results: list[str] = []

    for device_name in devices:
        for model_name in ("vision", "transformer"):
            for mode_name in ("inference", "training"):
                baseline = _run_workload(
                    model_name,
                    mode_name,
                    device_name,
                    profiled=False,
                    benchmark=False,
                )
                output_path = REPORT_DIR / (
                    f"{model_name}_{mode_name}_{device_name}_size.xlsx"
                )
                measured = _run_workload(
                    model_name,
                    mode_name,
                    device_name,
                    profiled=True,
                    benchmark=False,
                    output_path=output_path,
                )
                for actual, expected in zip(
                    measured["outputs"], baseline["outputs"]
                ):
                    torch.testing.assert_close(
                        torch.tensor(actual),
                        torch.tensor(expected),
                        rtol=1e-5,
                        atol=1e-5,
                    )
                if mode_name == "training":
                    assert len(measured["parameters"]) == len(
                        baseline["parameters"]
                    )
                    for actual, expected in zip(
                        measured["parameters"],
                        baseline["parameters"],
                    ):
                        torch.testing.assert_close(
                            torch.tensor(actual),
                            torch.tensor(expected),
                            rtol=1e-5,
                            atol=1e-5,
                        )
                _validate_report(output_path, model_name, mode_name)
                validation_results.append(
                    f"{model_name}-{mode_name}-{device_name}"
                )

                benchmark_baseline = _run_workload(
                    model_name,
                    mode_name,
                    device_name,
                    profiled=False,
                    benchmark=True,
                )
                benchmark_profiled = _run_workload(
                    model_name,
                    mode_name,
                    device_name,
                    profiled=True,
                    benchmark=True,
                    output_path=REPORT_DIR / (
                        f"benchmark_{model_name}_{mode_name}_{device_name}.xlsx"
                    ),
                )
                overhead = (
                    benchmark_profiled["elapsed_seconds"]
                    / benchmark_baseline["elapsed_seconds"]
                    - 1
                ) * 100
                benchmark_results.append(
                    f"{model_name}-{mode_name}-{device_name}: "
                    "mode=continuous_collection, "
                    f"rounds={benchmark_profiled['rounds']}, "
                    f"steps={benchmark_profiled['steps_per_round']}, "
                    f"baseline_median={benchmark_baseline['elapsed_seconds']:.6f}s, "
                    f"profiled_median={benchmark_profiled['elapsed_seconds']:.6f}s, "
                    f"overhead={overhead:.1f}%, "
                    f"collector_bytes="
                    f"{benchmark_profiled['collector_python_bytes']}"
                )

    print("REAL MODEL TEST PASS")
    print("validated:", ", ".join(validation_results))
    print("benchmarks:")
    for result in benchmark_results:
        print(result)
    print(f"reports: {REPORT_DIR}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--child", action="store_true")
    args = parser.parse_args()
    if args.child:
        _child_run()
    else:
        _run_validation()


if __name__ == "__main__":
    main()
