from __future__ import annotations

import os
import subprocess
import sys
import tempfile
import textwrap
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]


def _environment(**updates: str) -> dict[str, str]:
    environment = {
        key: value
        for key, value in os.environ.items()
        if not key.startswith("ATEN_SIZE_PROFILER_")
    }
    environment["PYTHONPATH"] = str(ROOT)
    environment.update(updates)
    return environment


def _run(name: str, source: str, **updates: str) -> subprocess.CompletedProcess[str]:
    result = subprocess.run(
        [sys.executable, "-c", textwrap.dedent(source)],
        cwd=ROOT,
        env=_environment(**updates),
        text=True,
        capture_output=True,
        timeout=240,
    )
    assert result.returncode == 0, (
        f"{name} failed with code {result.returncode}\n"
        f"stdout:\n{result.stdout}\nstderr:\n{result.stderr}"
    )
    return result


# R8: every schema out argument is excluded, not only names starting with
# "out"; in-place "self" arguments are still inputs.
OUT_ARGUMENT_SOURCE = """
import torch

from aten_size_profiler import AtenSizeMode

value = torch.arange(12.0).reshape(3, 4)
with AtenSizeMode() as mode:
    values = torch.empty(3, 2)
    indices = torch.empty(3, 2, dtype=torch.long)
    torch.topk(value, 2, out=(values, indices))
    max_values = torch.empty(3)
    max_indices = torch.empty(3, dtype=torch.long)
    torch.max(value, 1, out=(max_values, max_indices))
    sorted_values = torch.empty(3, 4)
    sorted_indices = torch.empty(3, 4, dtype=torch.long)
    torch.sort(value, out=(sorted_values, sorted_indices))
    added = torch.empty(3, 4)
    torch.add(value, value, out=added)
    in_place = value.clone()
    in_place.add_(1)

by_name = {}
for call in mode.calls:
    by_name.setdefault(call.op_name, []).append(call)
for op_name in (
    "aten.topk.values",
    "aten.max.dim_max",
    "aten.sort.values",
):
    paths = [tensor.path for tensor in by_name[op_name][0].inputs]
    assert paths == ["args[0]"], (op_name, paths)
add_paths = [tensor.path for tensor in by_name["aten.add.out"][0].inputs]
assert add_paths == ["args[0]", "args[1]"], add_paths
in_place_paths = [
    tensor.path for tensor in by_name["aten.add_.Tensor"][0].inputs
]
assert in_place_paths == ["args[0]"], in_place_paths
"""


# R9: dtype, device, memory_format and layout arguments are recorded as
# scalar_args, so casts to different dtypes are not merged in unique_shapes.
TORCH_ENUM_SCALAR_SOURCE = """
import json
from pathlib import Path

import torch
from openpyxl import load_workbook

from aten_size_profiler import AtenSizeMode, export_collector

output_path = Path(r"{output_path}")
value = torch.ones(2, 3)
with AtenSizeMode(record_scalar_args=True) as mode:
    value.to(torch.float16)
    value.to(torch.bfloat16)
    value.contiguous(memory_format=torch.contiguous_format).clone(
        memory_format=torch.preserve_format
    )
    torch.empty(2, layout=torch.strided, device="cpu")

copies = [call for call in mode.calls if call.op_name == "aten._to_copy.default"]
assert len(copies) == 2
copy_scalars = [dict(call.scalar_args) for call in copies]
assert copy_scalars[0]["kwargs['dtype']"] == json.dumps("torch.float16")
assert copy_scalars[1]["kwargs['dtype']"] == json.dumps("torch.bfloat16")

clone = next(call for call in mode.calls if call.op_name == "aten.clone.default")
assert dict(clone.scalar_args)["kwargs['memory_format']"] == json.dumps(
    "torch.preserve_format"
)
empty = next(
    call for call in mode.calls if call.op_name == "aten.empty.memory_format"
)
empty_scalars = dict(empty.scalar_args)
assert empty_scalars["kwargs['layout']"] == json.dumps("torch.strided")
assert empty_scalars["kwargs['device']"] == json.dumps("cpu")

export_collector(mode.collector, output_path)
workbook = load_workbook(output_path, read_only=True, data_only=True)
unique_copies = [
    row
    for row in workbook["unique_shapes"].iter_rows(min_row=2, values_only=True)
    if row[0] == "aten._to_copy.default"
]
assert len(unique_copies) == 2, unique_copies
workbook.close()
"""


# R7: a parameterized top-level module that is not the forward root triggers
# a RuntimeWarning and is listed in metadata; a parameterless loss module
# does not warn.
ROOT_MODULE_SOURCE = """
import json
import warnings
from pathlib import Path

import torch
from openpyxl import load_workbook
from torch import nn

import aten_size_profiler as profiler

class Preprocess(nn.Module):
    def __init__(self):
        super().__init__()
        self.scale = nn.Parameter(torch.ones(4))

    def forward(self, value):
        return value * self.scale

class Model(nn.Module):
    def __init__(self):
        super().__init__()
        self.linear = nn.Linear(4, 4)

    def forward(self, value):
        return self.linear(value)

preprocess = Preprocess()
model = Model()
loss_fn = nn.MSELoss()
value = torch.ones(2, 4)
with warnings.catch_warnings(record=True) as caught:
    warnings.simplefilter("always")
    for _ in range(3):
        loss_fn(model(preprocess(value)), value)

messages = [
    str(warning.message)
    for warning in caught
    if issubclass(warning.category, RuntimeWarning)
    and "aten_size_profiler" in str(warning.message)
]
assert len(messages) == 1, messages
assert "'Model'" in messages[0] and "'Preprocess'" in messages[0]

report = profiler.stop_profiler()
workbook = load_workbook(report, read_only=True, data_only=True)
metadata = dict(workbook["metadata"].iter_rows(min_row=2, values_only=True))
assert metadata["forward_root_module"] == "Preprocess"
assert json.loads(metadata["parameterized_non_root_modules"]) == ["Model"]
assert json.loads(metadata["other_top_level_modules"]) == ["MSELoss", "Model"]
workbook.close()
"""


LOSS_ONLY_SOURCE = """
import warnings

import torch
from torch import nn

import aten_size_profiler as profiler

model = nn.Linear(4, 4)
loss_fn = nn.MSELoss()
optimizer = torch.optim.SGD(model.parameters(), lr=0.1)
value = torch.ones(2, 4)
with warnings.catch_warnings(record=True) as caught:
    warnings.simplefilter("always")
    for _ in range(3):
        optimizer.zero_grad()
        loss_fn(model(value), value).backward()
        optimizer.step()
assert not [
    warning for warning in caught
    if "aten_size_profiler" in str(warning.message)
], [str(warning.message) for warning in caught]
assert profiler.get_iteration_status()["phase"] == "FINISHED"
assert profiler.stop_profiler(export=False) is None
"""


def run() -> None:
    with tempfile.TemporaryDirectory(prefix="aten-cross-review-round3-") as path:
        root = Path(path)
        _run(
            "out-arguments",
            OUT_ARGUMENT_SOURCE,
            ATEN_SIZE_PROFILER_ENABLED="false",
        )
        _run(
            "torch-enum-scalar-args",
            TORCH_ENUM_SCALAR_SOURCE.format(output_path=root / "enum.xlsx"),
            ATEN_SIZE_PROFILER_ENABLED="false",
        )
        _run(
            "root-module-warning",
            ROOT_MODULE_SOURCE,
            ATEN_SIZE_PROFILER_ENABLED="true",
            ATEN_SIZE_PROFILER_ITERATION_MODE="forward",
            ATEN_SIZE_PROFILER_TARGET_ITER="1",
            ATEN_SIZE_PROFILER_OUTPUT_DIR=str(root),
            ATEN_SIZE_PROFILER_OUTPUT_FILENAME="root.xlsx",
        )
        _run(
            "loss-module-no-warning",
            LOSS_ONLY_SOURCE,
            ATEN_SIZE_PROFILER_ENABLED="true",
            ATEN_SIZE_PROFILER_ITERATION_MODE="optimizer_step",
            ATEN_SIZE_PROFILER_TARGET_ITER="1",
        )

    print("CROSS REVIEW ROUND 3 TEST PASS")


if __name__ == "__main__":
    run()
