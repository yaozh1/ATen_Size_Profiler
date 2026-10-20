from __future__ import annotations

import os
import subprocess
import sys
import textwrap
from pathlib import Path

import torch
import yaml
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = (
    ROOT / "aten_size_reports" / "stage3_validation" / f"run_{os.getpid()}"
)
EXPECTED_SHEETS = [
    "summary",
    "unique_shapes",
    "detail",
    "input_tensors",
    "metadata",
]


def _child_environment(
    output_dir: Path,
    output_filename: str,
    *,
    enabled: bool,
    device: str = "cpu",
    config_path: Path | None = None,
    extra_environment: dict[str, str] | None = None,
) -> dict[str, str]:
    environment = dict(os.environ)
    for key in list(environment):
        if key.startswith("ATEN_SIZE_PROFILER_"):
            environment.pop(key)
    environment["PYTHONPATH"] = str(ROOT)
    environment["ATEN_SIZE_PROFILER_ENABLED"] = str(enabled).lower()
    environment["ATEN_SIZE_PROFILER_OUTPUT_DIR"] = str(output_dir)
    environment["ATEN_SIZE_PROFILER_OUTPUT_FILENAME"] = output_filename
    environment["TEST_DEVICE"] = device
    if config_path is not None:
        environment["ATEN_SIZE_PROFILER_CONFIG"] = str(config_path)
        environment.pop("ATEN_SIZE_PROFILER_OUTPUT_DIR", None)
        environment.pop("ATEN_SIZE_PROFILER_OUTPUT_FILENAME", None)
    if extra_environment is not None:
        environment.update(extra_environment)
    return environment


def _run_child(
    name: str,
    source: str,
    output_dir: Path,
    output_filename: str,
    *,
    enabled: bool = True,
    device: str = "cpu",
    config_path: Path | None = None,
    expect_success: bool = True,
    extra_environment: dict[str, str] | None = None,
) -> Path:
    result = subprocess.run(
        [sys.executable, "-c", textwrap.dedent(source)],
        cwd=ROOT,
        env=_child_environment(
            output_dir,
            output_filename,
            enabled=enabled,
            device=device,
            config_path=config_path,
            extra_environment=extra_environment,
        ),
        text=True,
        capture_output=True,
    )
    if expect_success:
        assert result.returncode == 0, (
            f"{name} failed with code {result.returncode}\n"
            f"stdout:\n{result.stdout}\nstderr:\n{result.stderr}"
        )
    else:
        assert result.returncode != 0, f"{name} unexpectedly succeeded"
    return output_dir / output_filename


def _rows(workbook, sheet_name: str):
    return list(workbook[sheet_name].iter_rows(min_row=2, values_only=True))


def _validate_report(path: Path, expected_mm_count: int | None = None) -> None:
    assert path.is_file(), f"report not found: {path}"
    workbook = load_workbook(path, read_only=True, data_only=True)
    assert workbook.sheetnames == EXPECTED_SHEETS
    summary_rows = _rows(workbook, "summary")
    detail_rows = _rows(workbook, "detail")
    metadata_rows = _rows(workbook, "metadata")
    metadata = {row[0]: row[1] for row in metadata_rows}
    assert detail_rows
    assert sum(row[1] for row in summary_rows) == len(detail_rows)
    assert metadata["call_count"] >= 0
    if expected_mm_count is not None:
        mm_rows = [row for row in summary_rows if row[0] == "aten.mm.default"]
        assert mm_rows and mm_rows[0][1] == expected_mm_count
    workbook.close()


def _run_auto_result_test(device: str, filename: str) -> Path:
    return _run_child(
        f"auto-{device}",
        """
        import os
        import torch

        torch.manual_seed(2026)
        device = torch.device(os.environ["TEST_DEVICE"])
        left = torch.randn(8, 16, device=device)
        right = torch.randn(16, 4, device=device)
        baseline = torch.relu(torch.mm(left, right))

        import aten_size_profiler as profiler

        assert profiler.is_installed()
        assert profiler.is_recording()
        measured = torch.relu(torch.mm(left, right))
        torch.testing.assert_close(measured, baseline)
        assert profiler.get_status()["export_count"] == 0
        """,
        REPORT_DIR,
        filename,
        device=device,
    )


def run() -> None:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)

    devices = ["cpu"]
    if torch.cuda.is_available():
        devices.append("cuda")
    for device in devices:
        report = _run_auto_result_test(
            device,
            f"auto_{device}_size.xlsx",
        )
        _validate_report(report)

    disabled_report = _run_child(
        "disabled",
        """
        import torch

        left = torch.ones(2, 2)
        right = torch.ones(2, 2)
        import aten_size_profiler as profiler
        assert profiler.is_installed()
        assert not profiler.is_recording()
        profiler.resume_profiler()
        assert not profiler.is_recording()
        torch.mm(left, right)
        assert profiler.get_collector().calls == []
        """,
        REPORT_DIR,
        "disabled_size.xlsx",
        enabled=False,
    )
    assert not disabled_report.exists()

    repeated_report = _run_child(
        "repeated-import",
        """
        import importlib
        import torch

        left = torch.ones(2, 2)
        right = torch.ones(2, 2)
        import aten_size_profiler as profiler
        importlib.reload(profiler)
        importlib.reload(profiler)
        assert profiler.get_status()["export_count"] == 0
        torch.mm(left, right)
        """,
        REPORT_DIR,
        "repeated_import_size.xlsx",
    )
    _validate_report(repeated_report, expected_mm_count=1)

    paused_report = _run_child(
        "pause-resume",
        """
        import torch

        left = torch.ones(2, 2)
        right = torch.ones(2, 2)
        import aten_size_profiler as profiler
        profiler.pause_profiler()
        torch.mm(left, right)
        assert profiler.get_collector().calls == []
        profiler.resume_profiler()
        torch.mm(left, right)
        report = profiler.stop_profiler()
        assert report is not None
        assert profiler.stop_profiler() == report
        assert profiler.get_status()["export_count"] == 1
        """,
        REPORT_DIR,
        "pause_resume_size.xlsx",
    )
    _validate_report(paused_report, expected_mm_count=1)

    exception_report = _run_child(
        "exception-exit",
        """
        import torch

        left = torch.ones(2, 2)
        right = torch.ones(2, 2)
        import aten_size_profiler
        torch.mm(left, right)
        raise RuntimeError("expected validation exception")
        """,
        REPORT_DIR,
        "exception_exit_size.xlsx",
        expect_success=False,
    )
    _validate_report(exception_report, expected_mm_count=1)

    yaml_path = REPORT_DIR / "stage3_config.yaml"
    yaml_path.write_text(
        yaml.safe_dump(
            {
                "enabled": False,
                "output_dir": str(REPORT_DIR),
                "output_filename": "yaml_config_size.xlsx",
                "record_stride": False,
                "record_stack": True,
                "iteration_mode": "manual",
                "target_iter": 7,
            },
            sort_keys=False,
        ),
        encoding="utf-8",
    )
    yaml_report = _run_child(
        "yaml-config",
        """
        import torch

        left = torch.ones(2, 2)
        right = torch.ones(2, 2)
        import aten_size_profiler
        torch.mm(left, right)
        """,
        REPORT_DIR,
        "ignored_by_yaml_size.xlsx",
        config_path=yaml_path,
    )
    yaml_report = REPORT_DIR / "yaml_config_size.xlsx"
    _validate_report(yaml_report, expected_mm_count=1)
    workbook = load_workbook(yaml_report, read_only=True, data_only=True)
    metadata = {
        row[0]: row[1]
        for row in _rows(workbook, "metadata")
    }
    detail_rows = _rows(workbook, "detail")
    assert metadata["record_stride"] is False
    assert metadata["record_stack"] is True
    assert metadata["configured_enabled"] is True
    assert metadata["target_iter"] == 7
    assert detail_rows[0][7] == "[null,null]"
    assert detail_rows[0][11]
    workbook.close()

    priority_path = REPORT_DIR / "priority_config.yaml"
    priority_path.write_text(
        yaml.safe_dump(
            {
                "enabled": False,
                "output_dir": str(REPORT_DIR / "yaml_output"),
                "output_filename": "yaml_priority_size.xlsx",
                "record_stride": True,
                "record_stack": False,
                "iteration_mode": "forward",
                "target_iter": 7,
            },
            sort_keys=False,
        ),
        encoding="utf-8",
    )
    _run_child(
        "config-priority",
        """
        import os
        import torch

        import aten_size_profiler as profiler

        config = profiler.load_config()
        assert config.enabled is True
        assert str(config.output_dir) == os.environ["ATEN_SIZE_PROFILER_OUTPUT_DIR"]
        assert config.output_filename == "env_priority_size.xlsx"
        assert config.record_stride is False
        assert config.record_stack is True
        assert config.iteration_mode == "optimizer_step"
        assert config.target_iter is None
        torch.mm(torch.ones(2, 2), torch.ones(2, 2))
        report = profiler.stop_profiler()
        assert report is not None
        """,
        REPORT_DIR,
        "ignored_by_priority_size.xlsx",
        config_path=priority_path,
        extra_environment={
            "ATEN_SIZE_PROFILER_OUTPUT_DIR": str(REPORT_DIR / "env_output"),
            "ATEN_SIZE_PROFILER_OUTPUT_FILENAME": "env_priority_size.xlsx",
            "ATEN_SIZE_PROFILER_RECORD_STRIDE": "false",
            "ATEN_SIZE_PROFILER_RECORD_STACK": "true",
            "ATEN_SIZE_PROFILER_ITERATION_MODE": "optimizer_step",
            "ATEN_SIZE_PROFILER_TARGET_ITER": "",
        },
    )
    priority_report = (
        REPORT_DIR / "env_output" / "env_priority_size.xlsx"
    )
    _validate_report(priority_report)

    default_path = REPORT_DIR / "default_config.yaml"
    default_path.write_text("{}\n", encoding="utf-8")
    _run_child(
        "config-default",
        """
        import aten_size_profiler as profiler

        assert profiler.load_config().iteration_mode == "forward"
        assert profiler.stop_profiler(export=False) is None
        """,
        REPORT_DIR,
        "ignored_by_default_size.xlsx",
        config_path=default_path,
    )

    no_export_report = _run_child(
        "stop-without-export",
        """
        import torch

        import aten_size_profiler as profiler

        torch.mm(torch.ones(2, 2), torch.ones(2, 2))
        assert profiler.stop_profiler(export=False) is None
        assert profiler.get_status()["export_count"] == 0
        """,
        REPORT_DIR,
        "stop_without_export_size.xlsx",
    )
    assert not no_export_report.exists()
    print("AUTO TEST PASS")
    print(f"reports: {REPORT_DIR}")


if __name__ == "__main__":
    run()
