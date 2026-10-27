"""Checks for record_memory_format: classification, config and report columns."""
from __future__ import annotations

import json
import os
import subprocess
import sys
import tempfile
import textwrap
import warnings
from pathlib import Path

import torch
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ["ATEN_SIZE_PROFILER_ENABLED"] = "false"

from aten_size_profiler import (
    AtenCollector,
    AtenSizeMode,
    ProfilerConfig,
    TensorInfo,
    collect_input_tensors,
    export_collector,
    load_config,
)
from aten_size_profiler.collector import classify_memory_layout


ENV_NAME = "ATEN_SIZE_PROFILER_RECORD_MEMORY_FORMAT"
SHEETS = ["summary", "unique_shapes", "detail", "input_tensors", "metadata"]
# Column layout of v0.2.1; record_memory_format=false must keep it exactly.
V021_HEADERS = {
    "summary": [
        "op_name", "call_count", "unique_shape_count",
        "input_tensor_count_total", "input_logical_bytes_total",
    ],
    "unique_shapes": [
        "op_name", "input_shapes", "input_dtypes", "input_devices",
        "input_strides", "count", "op_ratio", "overall_ratio", "scalar_args",
    ],
    "detail": [
        "call_id", "op_name", "input_count", "input_paths", "input_shapes",
        "input_dtypes", "input_devices", "input_strides", "input_numels",
        "input_element_sizes", "input_logical_bytes", "stack", "iteration",
        "rank", "scalar_args",
    ],
    "input_tensors": [
        "call_id", "op_name", "input_index", "input_path", "shape", "dtype",
        "device", "stride", "numel", "element_size", "logical_bytes",
        "iteration", "rank",
    ],
}
LAYOUT_COLUMNS = {
    "unique_shapes": "input_memory_layouts",
    "detail": "input_memory_layouts",
    "input_tensors": "memory_layout",
}


def _devices() -> list[torch.device]:
    devices = [torch.device("cpu")]
    if torch.cuda.is_available():
        devices.append(torch.device("cuda"))
    return devices


def _read(path: Path) -> dict[str, list[list]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == SHEETS, workbook.sheetnames
        report = {}
        for name in SHEETS:
            header, *rows = workbook[name].iter_rows(values_only=True)
            # Empty trailing cells (e.g. iteration=None) are not stored, so
            # read-only rows can be shorter than the header.
            width = len(header)
            report[name] = [list(header)] + [
                list(row) + [None] * (width - len(row)) for row in rows
            ]
        return report
    finally:
        workbook.close()


def _records(report: dict[str, list[list]], sheet: str) -> list[dict]:
    header, *rows = report[sheet]
    return [dict(zip(header, row)) for row in rows]


def _metadata(report: dict[str, list[list]]) -> dict:
    return {row[0]: row[1] for row in report["metadata"][1:]}


def _assert_headers(report: dict[str, list[list]], with_layout: bool) -> None:
    for sheet, header in V021_HEADERS.items():
        expected = list(header)
        if with_layout and sheet in LAYOUT_COLUMNS:
            expected.append(LAYOUT_COLUMNS[sheet])
        assert report[sheet][0] == expected, (sheet, report[sheet][0])


def check_classification() -> None:
    for device in _devices():
        x4 = torch.randn(2, 3, 4, 5, device=device)
        x5 = torch.randn(2, 3, 4, 5, 6, device=device)
        cases = [
            ("4D contiguous", x4, "NCHW"),
            ("4D N=1", torch.randn(1, 3, 4, 5, device=device), "NCHW"),
            ("4D channels_last",
             x4.contiguous(memory_format=torch.channels_last), "NHWC"),
            ("4D C=1", torch.randn(2, 1, 4, 5, device=device), "ambiguous"),
            ("4D H=W=1", torch.randn(2, 3, 1, 1, device=device), "ambiguous"),
            # Zero-element tensors are always contiguous in PyTorch.
            ("4D empty channels_last",
             torch.empty(0, 3, 4, 5, device=device).contiguous(
                 memory_format=torch.channels_last),
             "ambiguous"),
            ("4D transposed", x4.transpose(2, 3), "non_contiguous"),
            ("4D sliced", x4[:, :, ::2], "non_contiguous"),
            ("4D expanded",
             torch.randn(1, 3, 1, 1, device=device).expand(2, 3, 4, 5),
             "non_contiguous"),
            ("5D contiguous", x5, "NCDHW"),
            ("5D channels_last_3d",
             x5.contiguous(memory_format=torch.channels_last_3d), "NDHWC"),
            ("5D C=1", torch.randn(2, 1, 4, 5, 6, device=device), "ambiguous"),
            ("5D D=H=W=1",
             torch.randn(2, 3, 1, 1, 1, device=device), "ambiguous"),
            ("5D transposed", x5.transpose(1, 2), "non_contiguous"),
            ("0D", torch.tensor(1.0, device=device), "contiguous"),
            ("1D", torch.randn(8, device=device), "contiguous"),
            ("1D strided", torch.randn(8, device=device)[::2], "non_contiguous"),
            ("2D", torch.randn(3, 4, device=device), "contiguous"),
            ("2D transposed", torch.randn(3, 4, device=device).t(),
             "non_contiguous"),
            ("3D permuted", torch.randn(2, 3, 4, device=device).permute(0, 2, 1),
             "non_contiguous"),
            ("6D", torch.randn(1, 2, 1, 2, 1, 2, device=device), "contiguous"),
        ]
        for name, tensor, expected in cases:
            actual = classify_memory_layout(tensor)
            assert actual == expected, (device, name, actual, expected)

    meta = torch.empty(2, 3, 4, 5, device="meta").contiguous(
        memory_format=torch.channels_last
    )
    assert classify_memory_layout(meta) == "NHWC"

    # Sparse CSR has no is_contiguous(): the layout is None, and the tensor is
    # still recorded exactly as without record_memory_format.
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        csr = torch.randn(3, 3).to_sparse_csr()
    assert classify_memory_layout(csr) is None
    plain = collect_input_tensors((csr,), {}, record_stride=False)
    layout = collect_input_tensors(
        (csr,), {}, record_stride=False, record_memory_format=True,
    )
    assert len(plain) == len(layout) == 1
    assert layout[0].memory_layout is None
    assert plain[0].memory_layout is None
    print("classification: PASS")


def check_config(work_dir: Path) -> None:
    assert ProfilerConfig().record_memory_format is False
    assert load_config({}, work_dir).record_memory_format is False
    # Appended as the last field: v0.2.1 positional construction still works.
    positional = ProfilerConfig(
        True, Path("out"), "r.xlsx", False, True, False, 3, "optimizer_step",
        None,
    )
    assert positional.record_stride is False
    assert positional.target_iter == 3
    assert positional.iteration_mode == "optimizer_step"
    assert positional.record_memory_format is False

    for value in ("true", "1", "yes", "on", "TRUE"):
        assert load_config({ENV_NAME: value}, work_dir).record_memory_format
    for value in ("false", "0", "no", "off"):
        assert not load_config({ENV_NAME: value}, work_dir).record_memory_format

    yaml_path = work_dir / "layout.yaml"
    environment = {"ATEN_SIZE_PROFILER_CONFIG": str(yaml_path)}
    yaml_path.write_text("record_memory_format: true\n", encoding="utf-8")
    assert load_config(environment, work_dir).record_memory_format is True
    yaml_path.write_text("record_memory_format: false\n", encoding="utf-8")
    assert load_config(environment, work_dir).record_memory_format is False
    yaml_path.write_text("record_memory_format: true\n", encoding="utf-8")
    overridden = load_config({**environment, ENV_NAME: "false"}, work_dir)
    assert overridden.record_memory_format is False

    yaml_path.write_text("record_memory_format: maybe\n", encoding="utf-8")
    for source in (environment, {ENV_NAME: "maybe"}):
        try:
            load_config(source, work_dir)
        except ValueError as error:
            assert "record_memory_format" in str(error), error
        else:
            raise AssertionError(f"accepted invalid value via {source}")
    print("config: PASS")


def _run_model(device: torch.device, **mode_options) -> AtenSizeMode:
    torch.manual_seed(2026)
    x = torch.randn(2, 3, 8, 8, device=device)
    weight = torch.randn(4, 3, 3, 3, device=device)
    volume = torch.randn(1, 2, 3, 4, 5, device=device)
    matrix = torch.randn(4, 6, device=device)
    with AtenSizeMode(**mode_options) as mode:
        for memory_format in (torch.contiguous_format, torch.channels_last):
            torch.nn.functional.conv2d(
                x.contiguous(memory_format=memory_format),
                weight.contiguous(memory_format=memory_format),
                padding=1,
            )
            torch.relu(x.contiguous(memory_format=memory_format))
        torch.relu(volume)
        torch.relu(volume.contiguous(memory_format=torch.channels_last_3d))
        torch.relu(x.transpose(2, 3))
        torch.mm(matrix, matrix.t())
        torch.add(matrix, 1)
    return mode


def check_collection_compatibility() -> None:
    for device in _devices():
        default = _run_model(device)
        off = _run_model(device, record_memory_format=False)
        on = _run_model(device, record_memory_format=True)
        assert default.record_memory_format is False
        for mode in (default, off):
            assert all(
                tensor.memory_layout is None
                for call in mode.calls for tensor in call.inputs
            )
        assert len(on.calls) == len(off.calls) > 0
        layouts = set()
        for call_on, call_off in zip(on.calls, off.calls):
            assert call_on.op_name == call_off.op_name
            assert call_on.scalar_args == call_off.scalar_args
            assert len(call_on.inputs) == len(call_off.inputs)
            for tensor_on, tensor_off in zip(call_on.inputs, call_off.inputs):
                assert tensor_on.memory_layout is not None
                layouts.add(tensor_on.memory_layout)
                restored = tensor_on.__class__(
                    **{**tensor_on.__dict__, "memory_layout": None}
                )
                assert restored == tensor_off, (tensor_on, tensor_off)
        assert {
            "NCHW", "NHWC", "NCDHW", "NDHWC", "non_contiguous", "contiguous",
        } <= layouts, layouts
        # The in-memory aggregates keep their v0.2.1 key.
        assert on.aggregates.keys() == off.aggregates.keys()
        for op_name, aggregate in on.aggregates.items():
            other = off.aggregates[op_name]
            assert aggregate.call_count == other.call_count
            assert aggregate.input_signature_counts == other.input_signature_counts
    print("collection compatibility: PASS")


def _strip_layout(report: dict[str, list[list]]) -> dict[str, list[list]]:
    """Drop the appended layout columns, leaving the v0.2.1 columns."""
    stripped = {}
    for sheet, header in V021_HEADERS.items():
        width = len(header)
        stripped[sheet] = [row[:width] for row in report[sheet]]
    return stripped


def check_excel_columns(work_dir: Path) -> None:
    for device in _devices():
        off_path = export_collector(
            _run_model(device).collector, work_dir / f"{device.type}_off.xlsx",
        )
        on_path = export_collector(
            _run_model(device, record_memory_format=True).collector,
            work_dir / f"{device.type}_on.xlsx",
            record_memory_format=True,
        )
        off, on = _read(off_path), _read(on_path)
        _assert_headers(off, with_layout=False)
        _assert_headers(on, with_layout=True)
        assert _metadata(off)["record_memory_format"] is False
        assert _metadata(on)["record_memory_format"] is True
        # Enabling the option only appends columns; every original cell,
        # including unique_shapes grouping and counts, stays the same.
        assert _strip_layout(on) == _strip_layout(off)

        input_rows = _records(on, "input_tensors")
        detail_rows = _records(on, "detail")
        for row in detail_rows:
            layouts = json.loads(row["input_memory_layouts"])
            assert len(layouts) == row["input_count"], row
            assert layouts == [
                tensor["memory_layout"]
                for tensor in input_rows if tensor["call_id"] == row["call_id"]
            ]
        conv_layouts = [
            json.loads(row["input_memory_layouts"])
            for row in detail_rows
            if row["op_name"] == "aten.convolution.default"
        ]
        assert conv_layouts == [["NCHW", "NCHW"], ["NHWC", "NHWC"]]
        relu_layouts = [
            json.loads(row["input_memory_layouts"])
            for row in detail_rows if row["op_name"] == "aten.relu.default"
        ]
        assert relu_layouts == [
            ["NCHW"], ["NHWC"], ["NCDHW"], ["NDHWC"], ["non_contiguous"],
        ]
        mm_layouts = [
            json.loads(row["input_memory_layouts"])
            for row in detail_rows if row["op_name"] == "aten.mm.default"
        ]
        assert mm_layouts == [["contiguous", "non_contiguous"]]
        # record_stride=true: different layouts have different strides, so
        # they stay separate rows without mixed values.
        unique_rows = _records(on, "unique_shapes")
        assert all(
            "mixed" not in json.loads(row["input_memory_layouts"])
            for row in unique_rows
        )
    print("excel columns: PASS")


def check_mixed_groups(work_dir: Path) -> None:
    for device in _devices():
        nchw = torch.randn(2, 3, 4, 5, device=device)
        nhwc = nchw.contiguous(memory_format=torch.channels_last)
        with AtenSizeMode(record_stride=False, record_memory_format=True) as mode:
            torch.relu(nchw)
            torch.relu(nhwc)
            torch.relu(nchw)
            torch.add(nchw, nchw)
            torch.add(nchw, nhwc)
            torch.mul(nhwc, nchw)
            torch.mul(nhwc, nchw)
        report = _read(export_collector(
            mode.collector, work_dir / f"{device.type}_mixed.xlsx",
            record_memory_format=True,
        ))
        groups = {
            row["op_name"]: row for row in _records(report, "unique_shapes")
        }
        assert len(groups) == len(_records(report, "unique_shapes")) == 3
        # The first relu call is NCHW; the group must not report only it.
        assert groups["aten.relu.default"]["count"] == 3
        assert json.loads(
            groups["aten.relu.default"]["input_memory_layouts"]
        ) == ["mixed"]
        # Inputs are merged position by position.
        assert json.loads(
            groups["aten.add.Tensor"]["input_memory_layouts"]
        ) == ["NCHW", "mixed"]
        assert json.loads(
            groups["aten.mul.Tensor"]["input_memory_layouts"]
        ) == ["NHWC", "NCHW"]
        assert groups["aten.relu.default"]["input_strides"] == "[null]"
        detail = [
            json.loads(row["input_memory_layouts"])
            for row in _records(report, "detail")
        ]
        assert detail == [
            ["NCHW"], ["NHWC"], ["NCHW"], ["NCHW", "NCHW"], ["NCHW", "NHWC"],
            ["NHWC", "NCHW"], ["NHWC", "NCHW"],
        ]
        summary = {
            row["op_name"]: row for row in _records(report, "summary")
        }
        assert summary["aten.relu.default"]["unique_shape_count"] == 1
    print("mixed groups: PASS")


def _tensor_info(memory_layout: str | None) -> TensorInfo:
    return TensorInfo(
        path="args[0]", shape=(2,), dtype="torch.float32", device="cpu",
        stride=(1,), numel=2, element_size=4, logical_bytes=8,
        memory_layout=memory_layout,
    )


def check_export_argument(work_dir: Path) -> None:
    unset = AtenCollector()
    unset.record("aten.relu.default", (_tensor_info(None),))
    report = _read(export_collector(unset, work_dir / "unset.xlsx"))
    _assert_headers(report, with_layout=False)
    assert _metadata(report)["record_memory_format"] is False

    report = _read(export_collector(
        unset, work_dir / "unset_forced.xlsx", record_memory_format=True,
    ))
    _assert_headers(report, with_layout=True)
    assert report["input_tensors"][1][-1] is None
    assert report["detail"][1][-1] == "[null]"

    recorded = AtenCollector()
    recorded.record("aten.relu.default", (_tensor_info("contiguous"),))
    recorded.record("aten.relu.default", (_tensor_info(None),))
    report = _read(export_collector(recorded, work_dir / "recorded.xlsx"))
    _assert_headers(report, with_layout=True)
    assert _metadata(report)["record_memory_format"] is True
    assert json.loads(report["unique_shapes"][1][-1]) == ["mixed"]

    report = _read(export_collector(
        recorded, work_dir / "recorded_off.xlsx", record_memory_format=False,
    ))
    _assert_headers(report, with_layout=False)
    assert _metadata(report)["record_memory_format"] is False
    print("export argument: PASS")


AUTO_SOURCE = """
import torch

device = torch.device({device!r})
model = torch.nn.Conv2d(3, 4, 3, padding=1).to(device)
model = model.to(memory_format=torch.channels_last)
inputs = torch.randn(2, 3, 8, 8, device=device)
inputs = inputs.contiguous(memory_format=torch.channels_last)

import aten_size_profiler as profiler

for _ in range(3):
    model(inputs)
assert profiler.stop_profiler() is not None
"""


def _run_auto(
    name: str,
    work_dir: Path,
    device: str,
    **environment_updates: str,
) -> dict[str, list[list]]:
    environment = {
        key: value
        for key, value in os.environ.items()
        if not key.startswith("ATEN_SIZE_PROFILER_")
    }
    environment.update(
        PYTHONPATH=str(ROOT),
        ATEN_SIZE_PROFILER_OUTPUT_DIR=str(work_dir),
        ATEN_SIZE_PROFILER_OUTPUT_FILENAME=f"{name}.xlsx",
        ATEN_SIZE_PROFILER_TARGET_ITER="1",
        ATEN_SIZE_PROFILER_ITERATION_MODE="forward",
    )
    environment.update(environment_updates)
    result = subprocess.run(
        [sys.executable, "-c", textwrap.dedent(AUTO_SOURCE.format(device=device))],
        cwd=work_dir,
        env=environment,
        text=True,
        capture_output=True,
        timeout=240,
    )
    assert result.returncode == 0, (
        f"{name} failed with code {result.returncode}\n"
        f"stdout:\n{result.stdout}\nstderr:\n{result.stderr}"
    )
    return _read(work_dir / f"{name}.xlsx")


def check_auto_enable(work_dir: Path) -> None:
    yaml_path = work_dir / "auto_layout.yaml"
    yaml_path.write_text("record_memory_format: true\n", encoding="utf-8")
    for device in _devices():
        default = _run_auto(f"auto_default_{device.type}", work_dir, device.type)
        from_env = _run_auto(
            f"auto_env_{device.type}", work_dir, device.type,
            **{ENV_NAME: "true"},
        )
        from_yaml = _run_auto(
            f"auto_yaml_{device.type}", work_dir, device.type,
            ATEN_SIZE_PROFILER_CONFIG=str(yaml_path),
        )
        env_off = _run_auto(
            f"auto_env_off_{device.type}", work_dir, device.type,
            ATEN_SIZE_PROFILER_CONFIG=str(yaml_path), **{ENV_NAME: "false"},
        )
        for report in (default, env_off):
            _assert_headers(report, with_layout=False)
            assert _metadata(report)["record_memory_format"] is False
        for report in (from_env, from_yaml):
            _assert_headers(report, with_layout=True)
            metadata = _metadata(report)
            assert metadata["record_memory_format"] is True
            assert metadata["iteration_complete"] is True
            conv = [
                json.loads(row["input_memory_layouts"])
                for row in _records(report, "detail")
                if row["op_name"] == "aten.convolution.default"
            ]
            assert conv == [["NHWC", "NHWC", "contiguous"]], conv
            assert {
                row["iteration"] for row in _records(report, "input_tensors")
            } == {1}
            assert _strip_layout(report)["summary"] == default["summary"]
    print("auto enable: PASS")


def run() -> None:
    with tempfile.TemporaryDirectory(prefix="aten-memory-layout-") as path:
        work_dir = Path(path)
        check_classification()
        check_config(work_dir)
        check_collection_compatibility()
        check_excel_columns(work_dir)
        check_mixed_groups(work_dir)
        check_export_argument(work_dir)
        check_auto_enable(work_dir)
    print("MEMORY LAYOUT TEST PASS")


if __name__ == "__main__":
    run()
