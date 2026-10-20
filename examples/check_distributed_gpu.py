from __future__ import annotations

import argparse
import os
from collections import Counter
from pathlib import Path

import torch
import torch.distributed as dist
from openpyxl import load_workbook
from torch import nn


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = Path(
    os.environ.get(
        "ATEN_SIZE_PROFILER_OUTPUT_DIR",
        str(ROOT / "aten_size_reports" / "stage8_validation"),
    )
)
EXPECTED_SHEETS = [
    "summary",
    "unique_shapes",
    "detail",
    "input_tensors",
    "metadata",
]


class Model(nn.Module):
    def __init__(self) -> None:
        super().__init__()
        self.weight = nn.Parameter(torch.randn(4, 4))

    def forward(self, inputs: torch.Tensor) -> torch.Tensor:
        return torch.mm(inputs, self.weight)


def _rows(workbook, sheet_name: str):
    return list(workbook[sheet_name].iter_rows(min_row=2, values_only=True))


def _configure_profiler(mode: str) -> None:
    os.environ["ATEN_SIZE_PROFILER_ENABLED"] = "true"
    os.environ["ATEN_SIZE_PROFILER_OUTPUT_DIR"] = str(REPORT_DIR)
    os.environ["ATEN_SIZE_PROFILER_OUTPUT_FILENAME"] = (
        f"aten_size_profile_{mode}_size.xlsx"
    )
    os.environ["ATEN_SIZE_PROFILER_TARGET_ITER"] = "1"
    os.environ["ATEN_SIZE_PROFILER_ITERATION_MODE"] = mode
    os.environ["ATEN_SIZE_PROFILER_RECORD_STRIDE"] = "true"
    os.environ["ATEN_SIZE_PROFILER_RECORD_STACK"] = "false"


def _initialize_distributed() -> tuple[int, int, torch.device]:
    assert torch.cuda.is_available()
    local_rank = int(os.environ["LOCAL_RANK"])
    visible_devices = torch.cuda.device_count()
    assert 0 <= local_rank < visible_devices
    torch.cuda.set_device(local_rank)
    dist.init_process_group("nccl")
    rank = dist.get_rank()
    world_size = dist.get_world_size()
    assert rank == int(os.environ["RANK"])
    assert world_size == int(os.environ["WORLD_SIZE"])
    return rank, world_size, torch.device("cuda", local_rank)


def _install_collective_probe() -> tuple[list[str], dict[str, object]]:
    calls: list[str] = []
    originals: dict[str, object] = {}
    collective_names = (
        "all_reduce",
        "all_gather",
        "all_gather_into_tensor",
        "broadcast",
        "barrier",
        "reduce",
        "reduce_scatter",
        "all_to_all",
    )
    for name in collective_names:
        function = getattr(dist, name, None)
        if function is None:
            continue
        originals[name] = function

        def wrapper(*args, _name=name, **kwargs):
            calls.append(_name)
            return originals[_name](*args, **kwargs)

        setattr(dist, name, wrapper)
    return calls, originals


def _restore_collective_probe(originals: dict[str, object]) -> None:
    for name, function in originals.items():
        setattr(dist, name, function)


def _validate_report(
    output_path: Path,
    rank: int,
    world_size: int,
    device: torch.device,
    mode: str,
    target_rows: int,
    other_rank_rows: int,
    expected_outputs: int,
) -> None:
    assert output_path.is_file(), output_path
    workbook = load_workbook(output_path, read_only=True, data_only=True)
    assert workbook.sheetnames == EXPECTED_SHEETS
    rows = {
        sheet_name: _rows(workbook, sheet_name)
        for sheet_name in EXPECTED_SHEETS
    }

    summary_counts = {row[0]: row[1] for row in rows["summary"]}
    detail_counts = Counter(row[1] for row in rows["detail"])
    unique_counts: Counter[str] = Counter()
    for row in rows["unique_shapes"]:
        unique_counts[row[0]] += row[5]

    assert rows["detail"]
    assert rows["input_tensors"]
    assert detail_counts == summary_counts
    assert unique_counts == summary_counts
    assert sum(row[1] for row in rows["summary"]) == len(rows["detail"])
    assert sum(row[5] for row in rows["unique_shapes"]) == len(
        rows["detail"]
    )
    assert sum(row[3] for row in rows["summary"]) == len(
        rows["input_tensors"]
    )
    assert [row[0] for row in rows["detail"]] == list(
        range(len(rows["detail"]))
    )

    assert {row[12] for row in rows["detail"]} == {1}
    assert {row[11] for row in rows["input_tensors"]} == {1}
    assert {row[13] for row in rows["detail"]} == {rank}
    assert {row[12] for row in rows["input_tensors"]} == {rank}
    input_shapes = {row[4] for row in rows["input_tensors"]}
    assert f"[{target_rows},4]" in input_shapes
    assert f"[{other_rank_rows},4]" not in input_shapes

    operators = set(detail_counts)
    assert "aten.mm.default" in operators
    assert summary_counts["aten.mm.default"] >= expected_outputs
    if mode == "optimizer_step":
        assert "aten.t.default" in operators
        assert operators.intersection(
            {"aten.add_.Tensor", "aten._foreach_add_.List"}
        )

    metadata = dict(rows["metadata"])
    assert metadata["call_count"] == len(rows["detail"])
    assert metadata["rank"] == rank
    assert metadata["global_rank"] == rank
    assert metadata["world_size"] == world_size
    assert metadata["distributed_initialized"] is True
    assert metadata["target_iter"] == 1
    assert metadata["iteration_mode"] == mode
    assert str(device) in str(metadata["devices"])
    workbook.close()


def _build_inputs(rank: int, device: torch.device) -> list[torch.Tensor]:
    rows_by_iteration = (2 + rank * 3, 3 + rank * 3, 4 + rank * 3)
    return [
        torch.randn(rows, 4, device=device)
        for rows in rows_by_iteration
    ]


def _run_inference(
    rank: int,
    world_size: int,
    device: torch.device,
) -> Path:
    torch.manual_seed(2028 + rank)
    initial_model = Model()
    inputs = _build_inputs(rank, device)
    initial_state = {
        name: value.detach().clone()
        for name, value in initial_model.state_dict().items()
    }
    baseline = Model().to(device)
    baseline.load_state_dict(initial_state)
    baseline_outputs = [
        baseline(value).detach().clone()
        for value in inputs
    ]

    measured = Model().to(device)
    measured.load_state_dict(initial_state)
    import aten_size_profiler as profiler

    collective_calls, originals = _install_collective_probe()
    measured_outputs = [measured(value) for value in inputs]
    output_path = profiler.stop_profiler()
    _restore_collective_probe(originals)

    assert output_path is not None
    assert collective_calls == []
    for actual, expected in zip(measured_outputs, baseline_outputs):
        torch.testing.assert_close(actual, expected)

    target_rows = 3 + rank * 3
    other_rank_rows = 3 + (1 - rank) * 3
    _validate_report(
        output_path,
        rank,
        world_size,
        device,
        "forward",
        target_rows,
        other_rank_rows,
        expected_outputs=1,
    )
    return output_path


def _run_training(
    rank: int,
    world_size: int,
    device: torch.device,
) -> Path:
    torch.manual_seed(2030 + rank)
    initial_model = Model()
    initial_state = {
        name: value.detach().clone()
        for name, value in initial_model.state_dict().items()
    }
    inputs = _build_inputs(rank, device)

    baseline = Model().to(device)
    baseline.load_state_dict(initial_state)
    baseline_optimizer = torch.optim.SGD(baseline.parameters(), lr=0.01)
    baseline_outputs = []
    baseline_gradients = []
    for value in inputs:
        baseline_optimizer.zero_grad(set_to_none=True)
        output = baseline(value)
        output.square().mean().backward()
        baseline_gradients.append(
            [
                parameter.grad.detach().clone()
                for parameter in baseline.parameters()
            ]
        )
        baseline_optimizer.step()
        baseline_outputs.append(output.detach().clone())
    baseline_parameters = [
        parameter.detach().clone()
        for parameter in baseline.parameters()
    ]

    measured = Model().to(device)
    measured.load_state_dict(initial_state)
    measured_optimizer = torch.optim.SGD(measured.parameters(), lr=0.01)
    import aten_size_profiler as profiler

    collective_calls, originals = _install_collective_probe()
    measured_outputs = []
    measured_gradients = []
    for value in inputs:
        measured_optimizer.zero_grad(set_to_none=True)
        output = measured(value)
        output.square().mean().backward()
        measured_gradients.append(
            [
                parameter.grad.detach().clone()
                for parameter in measured.parameters()
            ]
        )
        measured_optimizer.step()
        measured_outputs.append(output.detach().clone())
    output_path = profiler.stop_profiler()
    _restore_collective_probe(originals)

    assert output_path is not None
    assert collective_calls == []
    for actual, expected in zip(measured_outputs, baseline_outputs):
        torch.testing.assert_close(actual, expected)
    for actual_step, expected_step in zip(
        measured_gradients,
        baseline_gradients,
    ):
        for actual, expected in zip(actual_step, expected_step):
            torch.testing.assert_close(actual, expected)
    for actual, expected in zip(measured.parameters(), baseline_parameters):
        torch.testing.assert_close(actual, expected)

    target_rows = 3 + rank * 3
    other_rank_rows = 3 + (1 - rank) * 3
    _validate_report(
        output_path,
        rank,
        world_size,
        device,
        "optimizer_step",
        target_rows,
        other_rank_rows,
        expected_outputs=1,
    )
    return output_path


def run(mode: str, probe_only: bool) -> None:
    if probe_only:
        rank, world_size, device = _initialize_distributed()
        print(
            f"RCCL INIT PASS rank={rank} world_size={world_size} "
            f"local_rank={device.index} device={device}",
            flush=True,
        )
        dist.destroy_process_group()
        return

    _configure_profiler(mode)
    rank, world_size, device = _initialize_distributed()
    assert world_size == 2
    output_path = (
        _run_inference(rank, world_size, device)
        if mode == "forward"
        else _run_training(rank, world_size, device)
    )
    print(
        f"GPU DISTRIBUTED {mode} RANK {rank} PASS: {output_path} "
        f"device={device}",
        flush=True,
    )
    dist.destroy_process_group()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--mode",
        choices=("forward", "optimizer_step"),
        default="forward",
    )
    parser.add_argument("--probe-only", action="store_true")
    args = parser.parse_args()
    run(args.mode, args.probe_only)


if __name__ == "__main__":
    main()
