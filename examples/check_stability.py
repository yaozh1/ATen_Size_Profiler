from __future__ import annotations

import os
import subprocess
import sys
import textwrap
from pathlib import Path

import torch
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = (
    ROOT / "aten_size_reports" / "stage6_validation" / f"run_{os.getpid()}"
)


TRAINING_SOURCE = """
import importlib
import os
import torch
from torch import nn

class Model(nn.Module):
    def __init__(self):
        super().__init__()
        self.weight = nn.Parameter(torch.randn(4, 4, device=device))
        self.relu = nn.ReLU()

    def forward(self, inputs):
        return self.relu(torch.mm(inputs, self.weight))

device = torch.device(os.environ["TEST_DEVICE"])
torch.manual_seed(2026)
model = Model()
baseline = Model()
baseline.load_state_dict(model.state_dict())
inputs = [torch.randn(rows, 4, device=device) for rows in (2, 3, 5)]
baseline_optimizer = torch.optim.SGD(baseline.parameters(), lr=0.01)
expected = []
for value in inputs:
    baseline_optimizer.zero_grad(set_to_none=True)
    output = baseline(value)
    loss = output.square().mean()
    loss.backward()
    expected.append((output.detach().clone(), loss.detach().clone(),
                     baseline.weight.grad.detach().clone()))
    baseline_optimizer.step()
del baseline_optimizer

optimizer = None
if os.environ["OPTIMIZER_ORDER"] == "before":
    optimizer = torch.optim.SGD(model.parameters(), lr=0.01)
import aten_size_profiler as profiler
importlib.reload(profiler)
if optimizer is None:
    optimizer = torch.optim.SGD(model.parameters(), lr=0.01)
actual = []
for value in inputs:
    optimizer.zero_grad(set_to_none=True)
    output = model(value)
    loss = output.square().mean()
    loss.backward()
    actual.append((output.detach().clone(), loss.detach().clone(),
                   model.weight.grad.detach().clone()))
    optimizer.step()
assert profiler.get_iteration_status()["phase"] == "FINISHED"
report = profiler.stop_profiler()
assert report is not None
assert profiler.stop_profiler() == report
assert profiler.get_status()["export_count"] == 1
for actual_values, expected_values in zip(actual, expected):
    for measured, reference in zip(actual_values, expected_values):
        torch.testing.assert_close(measured, reference)
torch.testing.assert_close(model.weight, baseline.weight)
"""


EXCEPTION_SOURCE = """
import os
import torch
from torch import nn

device = torch.device(os.environ["TEST_DEVICE"])
failure = RuntimeError("expected forward failure")

class Inner(nn.Module):
    def __init__(self):
        super().__init__()
        self.weight = nn.Parameter(torch.randn(4, 4, device=device))
        self.fail = True

    def forward(self, inputs):
        output = torch.mm(inputs, self.weight)
        if self.fail:
            raise failure
        return output

model = nn.Sequential(Inner())
optimizer = torch.optim.SGD(model.parameters(), lr=0.01)
import aten_size_profiler as profiler
target = int(os.environ["ATEN_SIZE_PROFILER_TARGET_ITER"])
training = os.environ["ATEN_SIZE_PROFILER_ITERATION_MODE"] == "optimizer_step"
if os.environ["UNCAUGHT"] == "true":
    model(torch.randn(7, 4, device=device))

def fail_forward():
    try:
        model(torch.randn(7, 4, device=device))
    except RuntimeError as error:
        assert error is failure
    else:
        raise AssertionError("forward did not fail")

def finish_iteration(output):
    if training:
        output.square().mean().backward()
        optimizer.step()
        optimizer.zero_grad(set_to_none=True)

if target == 1:
    fail_forward()
    assert profiler.get_iteration_status()["current_iter"] == 0
    assert profiler.get_iteration_status()["phase"] == "WAITING"
    model[0].fail = False
    finish_iteration(model(torch.randn(2, 4, device=device)))
    model[0].fail = True

fail_forward()
assert profiler.get_iteration_status()["phase"] == "WAITING"
assert profiler.get_iteration_status()["current_iter"] == target
assert not profiler.is_recording()
failed_count = len(profiler.get_collector().calls)
assert failed_count == 0
assert profiler.get_collector().aggregates == {}
profiler.resume_profiler()
assert not profiler.is_recording()
torch.ones(9, 4, device=device)
optimizer.step()
assert profiler.get_iteration_status()["current_iter"] == target
assert len(profiler.get_collector().calls) == failed_count

model[0].fail = False
value = torch.randn(3, 4, device=device)
reference = torch.mm(value, model[0].weight)
output = model(value)
torch.testing.assert_close(output, reference)
finish_iteration(output)
assert profiler.get_iteration_status()["phase"] == "FINISHED"
assert not profiler.is_recording()
finished_count = len(profiler.get_collector().calls)
assert finished_count > 0
if not training:
    assert sum(
        call.op_name == "aten.mm.default"
        for call in profiler.get_collector().calls
    ) == 1
model[0].fail = True
fail_forward()
assert profiler.get_iteration_status()["phase"] == "FINISHED"
assert profiler.get_iteration_status()["current_iter"] == target + 1
assert len(profiler.get_collector().calls) == finished_count
report = profiler.stop_profiler()
assert profiler.stop_profiler() == report
assert profiler.get_status()["export_count"] == 1
"""


def _run_child(
    name: str,
    source: str,
    device: str,
    iteration_mode: str,
    target_iter: int,
    *,
    optimizer_order: str = "before",
    uncaught: bool = False,
) -> Path:
    environment = {
        key: value
        for key, value in os.environ.items()
        if not key.startswith("ATEN_SIZE_PROFILER_")
    }
    filename = f"{name}_{device}_size.xlsx"
    environment.update(
        {
            "PYTHONPATH": str(ROOT),
            "ATEN_SIZE_PROFILER_ENABLED": "true",
            "ATEN_SIZE_PROFILER_OUTPUT_DIR": str(REPORT_DIR),
            "ATEN_SIZE_PROFILER_OUTPUT_FILENAME": filename,
            "ATEN_SIZE_PROFILER_TARGET_ITER": str(target_iter),
            "ATEN_SIZE_PROFILER_ITERATION_MODE": iteration_mode,
            "TEST_DEVICE": device,
            "OPTIMIZER_ORDER": optimizer_order,
            "UNCAUGHT": str(uncaught).lower(),
        }
    )
    result = subprocess.run(
        [sys.executable, "-c", textwrap.dedent(source)],
        cwd=ROOT,
        env=environment,
        text=True,
        capture_output=True,
        timeout=300,
    )
    if uncaught:
        assert result.returncode != 0
        assert "RuntimeError: expected forward failure" in result.stderr
    else:
        assert result.returncode == 0, (
            f"{name}-{device} failed with code {result.returncode}\n"
            f"stdout:\n{result.stdout}\nstderr:\n{result.stderr}"
        )
    return REPORT_DIR / filename


def _validate_report(
    path: Path,
    iteration: int,
    training: bool,
    *,
    expected_mm_count: int | None = None,
    empty: bool = False,
) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    assert workbook.sheetnames == [
        "summary", "unique_shapes", "detail", "input_tensors", "metadata",
    ]
    rows = {
        name: list(workbook[name].iter_rows(min_row=2, values_only=True))
        for name in workbook.sheetnames
    }
    metadata = dict(rows["metadata"])
    assert sum(row[1] for row in rows["summary"]) == len(rows["detail"])
    assert sum(row[5] for row in rows["unique_shapes"]) == len(rows["detail"])
    assert metadata["call_count"] == len(rows["detail"])
    if empty:
        assert not rows["detail"]
        assert not rows["input_tensors"]
        workbook.close()
        return
    assert rows["detail"] and rows["input_tensors"]
    assert {row[12] for row in rows["detail"]} == {iteration}
    assert {row[11] for row in rows["input_tensors"]} == {iteration}
    if expected_mm_count is not None:
        assert sum(row[1] == "aten.mm.default" for row in rows["detail"]) == (
            expected_mm_count
        )
    if training:
        operators = {row[1] for row in rows["detail"]}
        assert operators.intersection(
            {"aten.add_.Tensor", "aten._foreach_add_.List"}
        )
        assert "aten.mm.default" in operators
    workbook.close()


def run() -> None:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    devices = ["cpu"] + (["cuda"] if torch.cuda.is_available() else [])
    for device in devices:
        for order in ("before", "after"):
            path = _run_child(
                f"optimizer_{order}_import", TRAINING_SOURCE, device,
                "optimizer_step", 1, optimizer_order=order,
            )
            _validate_report(path, 1, True)
        for iteration_mode in ("forward", "optimizer_step"):
            for target in (0, 1):
                path = _run_child(
                    f"exception_{iteration_mode}_{target}", EXCEPTION_SOURCE,
                    device, iteration_mode, target,
                )
                _validate_report(
                    path,
                    target,
                    iteration_mode == "optimizer_step",
                    expected_mm_count=(
                        1 if iteration_mode == "forward" else None
                    ),
                )
            path = _run_child(
                f"uncaught_{iteration_mode}", EXCEPTION_SOURCE,
                device, iteration_mode, 0, uncaught=True,
            )
            _validate_report(path, 0, False, empty=True)
        print(f"STABILITY {device} PASS", flush=True)
    print("STABILITY TEST PASS")
    print(f"reports: {REPORT_DIR}")


if __name__ == "__main__":
    run()
