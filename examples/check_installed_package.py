from __future__ import annotations

import argparse
import importlib.metadata
import os
import site
import sys
from collections import Counter
from pathlib import Path

import torch
from openpyxl import load_workbook
from torch import nn


EXPECTED_SHEETS = [
    "summary",
    "unique_shapes",
    "detail",
    "input_tensors",
    "metadata",
]


class Model(nn.Module):
    def __init__(self) -> None:
        super().__init__()
        self.weight = nn.Parameter(torch.randn(4, 4))

    def forward(self, inputs: torch.Tensor) -> torch.Tensor:
        return torch.mm(inputs, self.weight)


def _rows(workbook, sheet_name: str):
    return list(workbook[sheet_name].iter_rows(min_row=2, values_only=True))


def _validate_report(path: Path, profiler) -> None:
    assert path.is_file(), path
    workbook = load_workbook(path, read_only=True, data_only=True)
    assert workbook.sheetnames == EXPECTED_SHEETS
    rows = {
        sheet_name: _rows(workbook, sheet_name)
        for sheet_name in EXPECTED_SHEETS
    }
    summary_counts = {row[0]: row[1] for row in rows["summary"]}
    detail_counts = Counter(row[1] for row in rows["detail"])
    unique_counts: Counter[str] = Counter()
    for row in rows["unique_shapes"]:
        unique_counts[row[0]] += row[5]
    assert rows["detail"]
    assert rows["input_tensors"]
    assert detail_counts == summary_counts
    assert unique_counts == summary_counts
    assert sum(row[3] for row in rows["summary"]) == len(
        rows["input_tensors"]
    )
    assert {row[12] for row in rows["detail"]} == {0}
    assert {row[11] for row in rows["input_tensors"]} == {0}
    metadata = dict(rows["metadata"])
    assert metadata["tool_version"] == profiler.__version__
    assert metadata["call_count"] == len(rows["detail"])
    assert metadata["target_iter"] == 0
    assert metadata["rank"] == 0
    workbook.close()


def _run_inference(device: torch.device, output_path: Path) -> None:
    torch.manual_seed(2040)
    baseline_model = Model().to(device)
    measured_model = Model().to(device)
    measured_model.load_state_dict(baseline_model.state_dict())
    inputs = [
        torch.randn(rows, 4, device=device)
        for rows in (2, 3)
    ]
    baseline_outputs = [
        baseline_model(value).detach().clone()
        for value in inputs
    ]

    import aten_size_profiler as profiler

    measured_outputs = [
        measured_model(value).detach().clone()
        for value in inputs
    ]
    for actual, expected in zip(measured_outputs, baseline_outputs):
        torch.testing.assert_close(actual, expected)
    report = profiler.stop_profiler()
    assert report == output_path
    _validate_report(report, profiler)


def _run_training(device: torch.device, output_path: Path) -> None:
    torch.manual_seed(2041)
    baseline_model = Model().to(device)
    measured_model = Model().to(device)
    measured_model.load_state_dict(baseline_model.state_dict())
    baseline_optimizer = torch.optim.SGD(baseline_model.parameters(), lr=0.01)
    measured_optimizer = torch.optim.SGD(measured_model.parameters(), lr=0.01)
    inputs = [
        torch.randn(rows, 4, device=device)
        for rows in (2, 3)
    ]

    baseline_outputs = []
    baseline_gradients = []
    for value in inputs:
        baseline_optimizer.zero_grad(set_to_none=True)
        output = baseline_model(value)
        output.square().mean().backward()
        baseline_gradients.append(baseline_model.weight.grad.detach().clone())
        baseline_outputs.append(output.detach().clone())
        baseline_optimizer.step()

    import aten_size_profiler as profiler

    measured_outputs = []
    measured_gradients = []
    for value in inputs:
        measured_optimizer.zero_grad(set_to_none=True)
        output = measured_model(value)
        output.square().mean().backward()
        measured_gradients.append(measured_model.weight.grad.detach().clone())
        measured_outputs.append(output.detach().clone())
        measured_optimizer.step()

    for actual, expected in zip(measured_outputs, baseline_outputs):
        torch.testing.assert_close(actual, expected)
    for actual, expected in zip(measured_gradients, baseline_gradients):
        torch.testing.assert_close(actual, expected)
    torch.testing.assert_close(measured_model.weight, baseline_model.weight)
    report = profiler.stop_profiler()
    assert report == output_path
    _validate_report(report, profiler)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--device", choices=("cpu", "cuda"), required=True)
    parser.add_argument(
        "--mode",
        choices=("forward", "optimizer_step"),
        required=True,
    )
    args = parser.parse_args()
    device = torch.device(args.device)
    assert args.device == "cpu" or torch.cuda.is_available()
    output_path = Path(
        os.environ["ATEN_SIZE_PROFILER_OUTPUT_DIR"]
    ) / os.environ["ATEN_SIZE_PROFILER_OUTPUT_FILENAME"]
    import aten_size_profiler as profiler

    module_path = Path(profiler.__file__).resolve()
    site_packages = [
        Path(path).resolve()
        for path in site.getsitepackages()
    ]
    venv_root = Path(sys.prefix).resolve()
    assert any(
        module_path == path or path in module_path.parents
        for path in site_packages
    ), module_path
    assert venv_root in module_path.parents, module_path
    installed_version = importlib.metadata.version("aten-size-profiler")
    assert installed_version == profiler.__version__
    if args.mode == "forward":
        _run_inference(device, output_path)
    else:
        _run_training(device, output_path)
    print(
        f"INSTALLED PACKAGE PASS device={device} mode={args.mode} "
        f"version={profiler.__version__} metadata_version={installed_version} "
        f"module={module_path} "
        f"report={output_path}",
        flush=True,
    )


if __name__ == "__main__":
    main()
