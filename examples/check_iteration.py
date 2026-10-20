from __future__ import annotations

import os
import subprocess
import sys
import textwrap
from collections import Counter
from pathlib import Path

import torch
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = (
    ROOT / "aten_size_reports" / "stage4_validation" / f"run_{os.getpid()}"
)
EXPECTED_SHEETS = [
    "summary",
    "unique_shapes",
    "detail",
    "input_tensors",
    "metadata",
]


def _environment(
    output_filename: str,
    device: str,
    iteration_mode: str,
) -> dict[str, str]:
    environment = dict(os.environ)
    for key in list(environment):
        if key.startswith("ATEN_SIZE_PROFILER_"):
            environment.pop(key)
    environment.update(
        {
            "PYTHONPATH": str(ROOT),
            "ATEN_SIZE_PROFILER_ENABLED": "true",
            "ATEN_SIZE_PROFILER_OUTPUT_DIR": str(REPORT_DIR),
            "ATEN_SIZE_PROFILER_OUTPUT_FILENAME": output_filename,
            "ATEN_SIZE_PROFILER_TARGET_ITER": "1",
            "ATEN_SIZE_PROFILER_ITERATION_MODE": iteration_mode,
            "TEST_DEVICE": device,
        }
    )
    return environment


def _run_child(
    name: str,
    source: str,
    output_filename: str,
    device: str,
    iteration_mode: str,
) -> Path:
    result = subprocess.run(
        [sys.executable, "-c", textwrap.dedent(source)],
        cwd=ROOT,
        env=_environment(output_filename, device, iteration_mode),
        text=True,
        capture_output=True,
    )
    assert result.returncode == 0, (
        f"{name} failed with code {result.returncode}\n"
        f"stdout:\n{result.stdout}\nstderr:\n{result.stderr}"
    )
    return REPORT_DIR / output_filename


def _rows(workbook, sheet_name: str):
    return list(workbook[sheet_name].iter_rows(min_row=2, values_only=True))


def _validate_report(
    path: Path,
    expected_iteration_mode: str,
    require_optimizer_update: bool,
) -> None:
    assert path.is_file(), f"report not found: {path}"
    workbook = load_workbook(path, read_only=True, data_only=True)
    assert workbook.sheetnames == EXPECTED_SHEETS
    detail_rows = _rows(workbook, "detail")
    input_rows = _rows(workbook, "input_tensors")
    metadata = {
        row[0]: row[1]
        for row in _rows(workbook, "metadata")
    }
    assert detail_rows
    assert input_rows
    assert metadata["iteration"] == 1
    assert metadata["target_iter"] == 1
    assert metadata["call_count"] == len(detail_rows)
    assert metadata["iteration_mode"] == expected_iteration_mode
    assert all(row[12] == 1 for row in detail_rows)
    assert all(row[11] == 1 for row in input_rows)
    summary_rows = _rows(workbook, "summary")
    unique_rows = _rows(workbook, "unique_shapes")
    summary_counts = {row[0]: row[1] for row in summary_rows}
    assert Counter(row[1] for row in detail_rows) == summary_counts
    unique_counts = Counter()
    for row in unique_rows:
        unique_counts[row[0]] += row[5]
    assert unique_counts == summary_counts
    assert Counter(row[0] for row in unique_rows) == {
        row[0]: row[2] for row in summary_rows
    }
    assert sum(row[3] for row in summary_rows) == len(input_rows)
    assert sum(row[4] for row in summary_rows) == sum(
        row[10] for row in detail_rows
    ) == sum(row[10] for row in input_rows)
    assert any(row[1] == "aten.mm.default" for row in detail_rows)
    if require_optimizer_update:
        assert any(
            row[1] in {
                "aten.add_.Tensor",
                "aten._foreach_add_.List",
            }
            for row in detail_rows
        )
    workbook.close()


def _run_inference(device: str) -> Path:
    return _run_child(
        f"inference-{device}",
        """
        import os
        import torch
        from torch import nn

        class Model(nn.Module):
            def __init__(self):
                super().__init__()
                self.weight = nn.Parameter(torch.randn(4, 4))

            def forward(self, inputs):
                return torch.mm(inputs, self.weight)

        import aten_size_profiler as profiler

        device = torch.device(os.environ["TEST_DEVICE"])
        model = Model().to(device)
        for rows in (2, 3, 4):
            model(torch.randn(rows, 4, device=device))
        assert profiler.get_iteration_status()["phase"] == "FINISHED"
        """,
        f"inference_{device}_size.xlsx",
        device,
        "forward",
    )


def _run_training(device: str) -> Path:
    return _run_child(
        f"training-{device}",
        """
        import os
        import torch
        from torch import nn

        import aten_size_profiler as profiler

        class Model(nn.Module):
            def __init__(self):
                super().__init__()
                self.weight = nn.Parameter(torch.randn(4, 4))

            def forward(self, inputs):
                return torch.mm(inputs, self.weight)

        device = torch.device(os.environ["TEST_DEVICE"])
        model = Model().to(device)
        optimizer = torch.optim.SGD(model.parameters(), lr=0.01)
        for rows in (2, 3, 4):
            inputs = torch.randn(rows, 4, device=device)
            optimizer.zero_grad(set_to_none=True)
            output = model(inputs)
            loss = output.square().mean()
            loss.backward()
            optimizer.step()
        assert profiler.get_iteration_status()["phase"] == "FINISHED"
        """,
        f"training_{device}_size.xlsx",
        device,
        "optimizer_step",
    )


def run() -> None:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    devices = ["cpu"]
    if torch.cuda.is_available():
        devices.append("cuda")

    for device in devices:
        inference_report = _run_inference(device)
        _validate_report(inference_report, "forward", False)
        training_report = _run_training(device)
        _validate_report(training_report, "optimizer_step", True)

    print("ITERATION TEST PASS")
    print(f"reports: {REPORT_DIR}")


if __name__ == "__main__":
    run()
