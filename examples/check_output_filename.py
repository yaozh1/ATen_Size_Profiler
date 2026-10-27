from __future__ import annotations

import os
import subprocess
import sys
import tempfile
import textwrap
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]

INVALID_FILENAMES = [
    "report.csv",
    "report.xls",
    "report.xlsm",
    "report",
    "report.xlsx.csv",
    ".xlsx",
    "",
    " report.xlsx",
    "report.xlsx ",
    "nested/report.xlsx",
    "nested\\report.xlsx",
    "../report.xlsx",
]


def _environment(**updates: str) -> dict[str, str]:
    environment = {
        key: value
        for key, value in os.environ.items()
        if not key.startswith("ATEN_SIZE_PROFILER_")
    }
    environment["PYTHONPATH"] = str(ROOT)
    environment.update(updates)
    return environment


def _run(
    name: str,
    source: str,
    *,
    cwd: Path,
    expect_success: bool = True,
    **updates: str,
) -> subprocess.CompletedProcess[str]:
    result = subprocess.run(
        [sys.executable, "-c", textwrap.dedent(source)],
        cwd=cwd,
        env=_environment(**updates),
        text=True,
        capture_output=True,
        timeout=240,
    )
    succeeded = result.returncode == 0
    assert succeeded == expect_success, (
        f"{name} exited with code {result.returncode}\n"
        f"stdout:\n{result.stdout}\nstderr:\n{result.stderr}"
    )
    return result


# load_config() validates YAML and environment values; start_profiler()
# validates direct ProfilerConfig objects before installing anything.
CONFIG_SOURCE = """
from pathlib import Path

import aten_size_profiler as profiler
from aten_size_profiler import ProfilerConfig, load_config

invalid = {invalid!r}
work_dir = Path({work_dir!r})

assert profiler.stop_profiler(export=False) is None

for name in ("report.xlsx", "Report.XLSX", "model_size.v2.xlsx"):
    config = load_config({{"ATEN_SIZE_PROFILER_OUTPUT_FILENAME": name}}, work_dir)
    assert config.output_filename == name
    assert config.output_path.name == name

for name in invalid:
    try:
        load_config({{"ATEN_SIZE_PROFILER_OUTPUT_FILENAME": name}}, work_dir)
    except ValueError as error:
        assert "output_filename" in str(error), error
    else:
        raise AssertionError(f"environment accepted {{name!r}}")

yaml_path = work_dir / "bad_filename.yaml"
for text in ("output_filename: report.csv\\n", "output_filename: 123\\n"):
    yaml_path.write_text(text, encoding="utf-8")
    try:
        load_config({{"ATEN_SIZE_PROFILER_CONFIG": str(yaml_path)}}, work_dir)
    except ValueError as error:
        assert "output_filename" in str(error), error
    else:
        raise AssertionError(f"YAML accepted {{text!r}}")

for name in invalid:
    try:
        profiler.start_profiler(ProfilerConfig(output_filename=name))
    except ValueError as error:
        assert "output_filename" in str(error), error
    else:
        raise AssertionError(f"start_profiler accepted {{name!r}}")
    assert not profiler.is_installed()
"""


# export_collector() is public API and rejects non-.xlsx destinations before
# creating directories or files.
EXPORT_SOURCE = """
from pathlib import Path

import torch
from openpyxl import load_workbook

from aten_size_profiler import AtenSizeMode, export_collector

work_dir = Path({work_dir!r})
with AtenSizeMode() as mode:
    torch.ones(2) + 1

for name in ("report.csv", "report", ".xlsx", "report.xlsx.csv"):
    target_dir = work_dir / "rejected" / name.replace(".", "_")
    try:
        export_collector(mode.collector, target_dir / name)
    except ValueError as error:
        assert "Excel output path" in str(error), error
    else:
        raise AssertionError(f"export_collector accepted {{name!r}}")
    assert not target_dir.exists(), target_dir

upper = export_collector(mode.collector, work_dir / "accepted" / "UPPER.XLSX")
assert upper.is_file()
workbook = load_workbook(upper, read_only=True)
assert "summary" in workbook.sheetnames
workbook.close()
"""


# A bad environment value fails at import time with a clear error and does
# not leave a report behind at exit.
IMPORT_SOURCE = """
import aten_size_profiler
"""


def run() -> None:
    with tempfile.TemporaryDirectory(prefix="aten-output-filename-") as path:
        work_dir = Path(path)
        _run(
            "config-validation",
            CONFIG_SOURCE.format(
                invalid=INVALID_FILENAMES,
                work_dir=str(work_dir),
            ),
            cwd=work_dir,
            ATEN_SIZE_PROFILER_ENABLED="false",
        )
        _run(
            "export-validation",
            EXPORT_SOURCE.format(work_dir=str(work_dir)),
            cwd=work_dir,
            ATEN_SIZE_PROFILER_ENABLED="false",
        )
        report_dir = work_dir / "import_reports"
        result = _run(
            "import-with-invalid-filename",
            IMPORT_SOURCE,
            cwd=work_dir,
            expect_success=False,
            ATEN_SIZE_PROFILER_ENABLED="true",
            ATEN_SIZE_PROFILER_OUTPUT_DIR=str(report_dir),
            ATEN_SIZE_PROFILER_OUTPUT_FILENAME="report.csv",
        )
        assert "ValueError" in result.stderr, result.stderr
        assert "output_filename" in result.stderr, result.stderr
        assert not report_dir.exists(), list(report_dir.iterdir())

    print("OUTPUT FILENAME TEST PASS")


if __name__ == "__main__":
    run()
