from __future__ import annotations

import os
import subprocess
import sys
import tempfile
import textwrap
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]


def _environment(**updates: str) -> dict[str, str]:
    environment = {
        key: value
        for key, value in os.environ.items()
        if not key.startswith("ATEN_SIZE_PROFILER_")
    }
    environment["PYTHONPATH"] = str(ROOT)
    environment.update(updates)
    return environment


def _run(
    name: str,
    source: str,
    **updates: str,
) -> subprocess.CompletedProcess[str]:
    result = subprocess.run(
        [sys.executable, "-c", COMMON_SOURCE + textwrap.dedent(source)],
        cwd=ROOT,
        env=_environment(**updates),
        text=True,
        capture_output=True,
        timeout=240,
    )
    assert result.returncode == 0, (
        f"{name} exited with code {result.returncode}\n"
        f"stdout:\n{result.stdout}\nstderr:\n{result.stderr}"
    )
    return result


# stop_profiler() must leave user TorchDispatchModes entered after the
# profiler where they are: never call their __exit__/__enter__ (a mode may
# not be reentrant, and CommDebugMode/FlopCounterMode reset their counts on
# enter), keep the stack order, and stop recording. Once such a mode exits,
# is_in_torch_dispatch_mode() may stay True (a documented limit); the flags
# Dynamo reads must be restored.
COMMON_SOURCE = """
import threading
import warnings

import torch
import torch.utils._python_dispatch as pd
from torch.utils._python_dispatch import TorchDispatchMode

import aten_size_profiler as profiler
from aten_size_profiler.config import ProfilerConfig

NO_MODE_FLAGS = (False, False, False)


def stack():
    return pd._get_current_dispatch_mode_stack()


def flags():
    return (
        pd._is_in_torch_dispatch_mode,
        pd._is_in_non_infra_torch_dispatch_mode,
        pd._is_in_any_mode_without_ignore_compile_internals,
    )


class UserMode(TorchDispatchMode):
    # Entering resets the op count; with once=True entering again raises.
    def __init__(self, once=False):
        super().__init__()
        self.once = once
        self.entered = False
        self.events = []
        self.ops = 0

    def __torch_dispatch__(self, func, types, args=(), kwargs=None):
        self.ops += 1
        return func(*args, **(kwargs or {}))

    def __enter__(self):
        self.events.append("enter")
        if self.once and self.entered:
            raise RuntimeError("UserMode cannot be entered twice")
        self.entered = True
        self.ops = 0
        return super().__enter__()

    def __exit__(self, *exc):
        self.events.append("exit")
        return super().__exit__(*exc)


def start():
    return profiler.start_profiler(ProfilerConfig(enabled=True))


def stop(export=False):
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        profiler.stop_profiler(export=export)
    return [
        str(warning.message)
        for warning in caught
        if issubclass(warning.category, RuntimeWarning)
    ]


def assert_stopped(mode):
    status = profiler.get_status()
    assert not profiler.is_installed() and not profiler.is_recording()
    assert not status["installed"] and not status["recording"], status
    assert not any(current is mode for current in stack()), stack()


def assert_no_mode_left():
    assert stack() == [], stack()
    assert flags()[1:] == NO_MODE_FLAGS[1:], flags()


def assert_passthrough(mode, *user_modes):
    # Two ops reach every user mode and none is recorded.
    recorded = len(mode.collector.calls)
    before = [user_mode.ops for user_mode in user_modes]
    result = torch.ones(2) + 1
    after = [user_mode.ops for user_mode in user_modes]
    assert after == [ops + 2 for ops in before], (before, after)
    assert result.tolist() == [2.0, 2.0]
    assert len(mode.collector.calls) == recorded


# Stop the profiler started by import (profiler on top: every flag is
# restored).
assert stop() == []
assert stack() == [] and flags() == NO_MODE_FLAGS
"""


MAIN_SOURCE = """
# Profiler on top: the regular __exit__ path.
mode = start()
torch.ones(2) + 1
assert mode.collector.calls
assert stop() == []
assert_stopped(mode)
assert stack() == [] and flags() == NO_MODE_FLAGS

# A reentrant and a non-reentrant mode above the profiler.
for once in (False, True):
    mode = start()
    user = UserMode(once=once)
    with user:
        torch.ones(2)
        ops = user.ops
        assert ops == 1 and mode.collector.calls
        assert stop() == []
        assert stack() == [user], stack()
        assert user.events == ["enter"] and user.ops == ops
        assert flags() == (True, True, True), flags()
        assert_stopped(mode)
        assert_passthrough(mode, user)
        assert stop() == []
        assert stack() == [user] and user.events == ["enter"]
    assert user.events == ["enter", "exit"]
    assert_no_mode_left()

# Nested modes above the profiler keep their order.
mode = start()
outer, inner = UserMode(), UserMode(once=True)
with outer, inner:
    torch.ones(2)
    assert stop() == []
    assert stack() == [outer, inner], stack()
    assert outer.events == inner.events == ["enter"]
    assert_stopped(mode)
    assert_passthrough(mode, outer, inner)
assert outer.events == inner.events == ["enter", "exit"]
assert_no_mode_left()

# Repeated stop/start with modes above: each start puts a new profiler
# mode on top and each stop removes only that one.
outer = UserMode(once=True)
first = start()
with outer:
    assert stop() == []
    for _ in range(3):
        mode = start()
        assert mode is not first and stack() == [outer, mode], stack()
        inner = UserMode(once=True)
        with inner:
            torch.ones(2)
            assert mode.collector.calls
            assert stop() == []
            assert stack() == [outer, inner], stack()
            assert_stopped(mode)
            assert_passthrough(mode, outer, inner)
        assert inner.events == ["enter", "exit"]
        assert stack() == [outer] and flags() == (True, True, True)
assert outer.events == ["enter", "exit"]
assert_no_mode_left()

# A user with block raising after stop, and stop in a finally block while
# the exception unwinds.
for stop_in_finally in (False, True):
    mode = start()
    user = UserMode(once=True)
    try:
        with user:
            try:
                torch.ones(2)
                if not stop_in_finally:
                    assert stop() == []
                raise ValueError("model failure")
            finally:
                if stop_in_finally:
                    assert stop() == []
                assert stack() == [user], stack()
    except ValueError:
        pass
    assert user.events == ["enter", "exit"]
    assert_stopped(mode)
    assert_no_mode_left()

# A torch mode that resets its counts on enter.
from torch.utils.flop_counter import FlopCounterMode

mode = start()
matrix = torch.randn(4, 4)
flop_counter = FlopCounterMode(display=False)
with flop_counter:
    torch.mm(matrix, matrix)
    assert flop_counter.get_total_flops() == 128
    assert stop() == []
    assert [type(m).__name__ for m in stack()] == ["_FlopCounterMode"]
    torch.mm(matrix, matrix)
    assert flop_counter.get_total_flops() == 256
assert_stopped(mode)
assert_no_mode_left()

# Model execution after stop: same results, still seen by the user mode,
# not recorded.
torch.manual_seed(0)
model = torch.nn.Linear(4, 3)
inputs = torch.randn(2, 4)
expected = model(inputs).detach()
mode = start()
user = UserMode(once=True)
with user:
    model(inputs).sum().backward()
    recorded = len(mode.collector.calls)
    assert recorded > 0
    assert stop() == []
    before = user.ops
    output = model(inputs)
    output.sum().backward()
    assert user.ops > before
    assert len(mode.collector.calls) == recorded
assert torch.equal(output.detach(), expected)
assert_no_mode_left()
print("MAIN OK")
"""

COMM_DEBUG_SOURCE = """
import torch.distributed as dist
from torch.distributed.tensor.debug import CommDebugMode
from torch.testing._internal.distributed.fake_pg import FakeStore

dist.init_process_group("fake", store=FakeStore(), rank=0, world_size=2)
mode = start()
comm_mode = CommDebugMode()
with comm_mode:
    dist.all_reduce(torch.ones(2))
    assert comm_mode.get_total_counts() == 1
    assert stop() == []
    assert stack() == [comm_mode], stack()
    assert comm_mode.get_total_counts() == 1
    dist.all_reduce(torch.ones(2))
    assert comm_mode.get_total_counts() == 2
assert_stopped(mode)
assert_no_mode_left()
dist.destroy_process_group()
print("COMM OK")
"""


# At exit stop_profiler() exports with a non-reentrant mode still entered
# after the profiler.
ATEXIT_SOURCE = """
from pathlib import Path

mode = profiler.start_profiler(
    ProfilerConfig(
        enabled=True,
        output_dir=Path({output_dir!r}),
        output_filename="atexit.xlsx",
    )
)
user = UserMode(once=True)
user.__enter__()
torch.ones(2) + 1
assert mode.collector.calls
print("ATEXIT OK")
"""


# Orderings stop_profiler() cannot fully undo report a RuntimeWarning, stop
# recording and leave the other modes alone.
LIMITS_SOURCE = """
# Started inside a user with block: the user __exit__ pops the top mode,
# which is the profiler, and leaves the user mode on the stack.
user = UserMode()
with user:
    mode = start()
assert stack() == [user], stack()
messages = stop()
assert len(messages) == 1 and "not on this thread" in messages[0], messages
assert_stopped(mode)
assert_passthrough(mode, user)
assert user.events == ["enter", "exit"]
assert flags() == NO_MODE_FLAGS
# The user mode already ran its __exit__; only the stack entry is left.
torch._C._pop_torch_dispatch_stack(None)
assert stack() == []

# Stopped from another thread: the mode stack is thread-local, so the
# profiler mode stays on the starting thread's stack but stops recording.
mode = start()
torch.ones(2)
results = []
thread = threading.Thread(target=lambda: results.append(stop()))
thread.start()
thread.join()
assert len(results[0]) == 1 and "not on this thread" in results[0][0]
assert stack() == [mode] and not profiler.is_recording()
recorded = len(mode.collector.calls)
torch.ones(2) + 1
assert len(mode.collector.calls) == recorded
print("LIMITS OK")
"""


def _call_count(path: Path) -> object:
    workbook = load_workbook(path, read_only=True)
    try:
        rows = workbook["metadata"].iter_rows(min_row=2, values_only=True)
        return {row[0]: row[1] for row in rows}["call_count"]
    finally:
        workbook.close()


def run() -> None:
    result = _run("main", MAIN_SOURCE, ATEN_SIZE_PROFILER_ENABLED="true")
    assert "MAIN OK" in result.stdout, result.stdout
    assert "RuntimeWarning" not in result.stderr, result.stderr

    result = _run("comm-debug", COMM_DEBUG_SOURCE)
    assert "COMM OK" in result.stdout, result.stdout
    assert "RuntimeWarning" not in result.stderr, result.stderr

    result = _run("limits", LIMITS_SOURCE)
    assert "LIMITS OK" in result.stdout, result.stdout

    with tempfile.TemporaryDirectory(prefix="aten-mode-stack-") as path:
        output_dir = Path(path)
        result = _run(
            "atexit",
            ATEXIT_SOURCE.format(output_dir=str(output_dir)),
        )
        assert "ATEXIT OK" in result.stdout, result.stdout
        assert "export at exit failed" not in result.stderr, result.stderr
        assert "RuntimeWarning" not in result.stderr, result.stderr
        call_count = _call_count(output_dir / "atexit.xlsx")
        assert isinstance(call_count, int) and call_count > 0, call_count
    print("MODE STACK TEST PASS")


if __name__ == "__main__":
    run()
