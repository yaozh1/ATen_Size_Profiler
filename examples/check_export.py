from __future__ import annotations

import os
import sys
from collections import Counter
from pathlib import Path

import torch
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
os.environ["ATEN_SIZE_PROFILER_ENABLED"] = "false"

from aten_size_profiler import AtenSizeMode, export_collector
from aten_size_profiler.exporter import (
    MAX_EXCEL_ROWS,
    _validate_row_limits,
)


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "aten_size_reports"
EXPECTED_SHEETS = [
    "summary",
    "unique_shapes",
    "detail",
    "input_tensors",
    "metadata",
]


def _data_rows(worksheet):
    return list(worksheet.iter_rows(min_row=2, values_only=True))


def _build_report(device: torch.device) -> Path:
    torch.manual_seed(2026)
    left = torch.randn(32, 128, device=device, requires_grad=True)
    right = torch.randn(128, 256, device=device, requires_grad=True)

    with AtenSizeMode() as mode:
        output = torch.relu(torch.mm(left, right))
        loss = output.square().mean()
        loss.backward()

    output_path = REPORT_DIR / f"aten_size_{device.type}_size.xlsx"
    export_collector(
        mode.collector,
        output_path,
        metadata={"validation_device": str(left.device)},
        rank=0,
    )
    return output_path


def _validate_report(output_path: Path) -> None:
    workbook = load_workbook(output_path, read_only=True, data_only=True)
    assert workbook.sheetnames == EXPECTED_SHEETS

    summary_rows = _data_rows(workbook["summary"])
    unique_rows = _data_rows(workbook["unique_shapes"])
    detail_rows = _data_rows(workbook["detail"])
    input_rows = _data_rows(workbook["input_tensors"])
    metadata_rows = _data_rows(workbook["metadata"])
    metadata = {row[0]: row[1] for row in metadata_rows}

    assert summary_rows
    assert detail_rows
    assert metadata_rows
    assert metadata["rank"] == 0
    assert metadata["iteration"] == "not configured"
    assert all(row[12] is None for row in detail_rows)
    assert all(row[11] is None for row in input_rows)
    assert sum(row[1] for row in summary_rows) == len(detail_rows)
    assert sum(row[5] for row in unique_rows) == len(detail_rows)
    assert sum(row[3] for row in summary_rows) == len(input_rows)

    summary_calls = {row[0]: row[1] for row in summary_rows}
    summary_unique = {row[0]: row[2] for row in summary_rows}
    detail_calls = Counter(row[1] for row in detail_rows)
    unique_shape_counts = Counter(row[0] for row in unique_rows)
    unique_call_counts = Counter()
    for row in unique_rows:
        unique_call_counts[row[0]] += row[5]
    assert dict(detail_calls) == summary_calls
    assert dict(unique_call_counts) == summary_calls
    assert dict(unique_shape_counts) == summary_unique

    detail_call_ids = [row[0] for row in detail_rows]
    assert detail_call_ids == list(range(len(detail_rows)))
    assert all(row[6] >= 0 for row in unique_rows)
    assert all(row[7] >= 0 for row in unique_rows)

    workbook.close()


def _validate_row_limit_guard() -> None:
    for sheet_name in EXPECTED_SHEETS:
        row_counts = {
            name: 1
            for name in EXPECTED_SHEETS
        }
        row_counts[sheet_name] = MAX_EXCEL_ROWS
        _validate_row_limits(row_counts)
        row_counts[sheet_name] = MAX_EXCEL_ROWS + 1
        try:
            _validate_row_limits(row_counts)
        except ValueError as error:
            assert sheet_name in str(error)
        else:
            raise AssertionError(
                f"expected row limit failure for {sheet_name}"
            )


def run() -> None:
    _validate_row_limit_guard()
    devices = [torch.device("cpu")]
    if torch.cuda.is_available():
        devices.append(torch.device("cuda"))

    output_paths = []
    for device in devices:
        output_path = _build_report(device)
        _validate_report(output_path)
        output_paths.append(output_path)

    for output_path in output_paths:
        print(f"validated: {output_path}")
    print("EXCEL TEST PASS")


if __name__ == "__main__":
    run()
