from __future__ import annotations

import json
import os
import subprocess
import sys
import tempfile
import textwrap
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SHEETS = [
    "summary",
    "unique_shapes",
    "detail",
    "input_tensors",
    "metadata",
]


def _environment(**updates: str) -> dict[str, str]:
    environment = {
        key: value
        for key, value in os.environ.items()
        if not key.startswith("ATEN_SIZE_PROFILER_")
    }
    environment["PYTHONPATH"] = str(ROOT)
    environment.update(updates)
    return environment


def _run(name: str, source: str, **updates: str) -> subprocess.CompletedProcess[str]:
    result = subprocess.run(
        [sys.executable, "-c", textwrap.dedent(source)],
        cwd=ROOT,
        env=_environment(**updates),
        text=True,
        capture_output=True,
        timeout=240,
    )
    assert result.returncode == 0, (
        f"{name} failed with code {result.returncode}\n"
        f"stdout:\n{result.stdout}\nstderr:\n{result.stderr}"
    )
    return result


def _run_file(
    name: str,
    source: str,
    work_dir: Path,
    **updates: str,
) -> subprocess.CompletedProcess[str]:
    script_path = work_dir / f"{name}.py"
    script_path.write_text(textwrap.dedent(source), encoding="utf-8")
    result = subprocess.run(
        [sys.executable, str(script_path)],
        cwd=ROOT,
        env=_environment(**updates),
        text=True,
        capture_output=True,
        timeout=240,
    )
    assert result.returncode == 0, (
        f"{name} failed with code {result.returncode}\n"
        f"stdout:\n{result.stdout}\nstderr:\n{result.stderr}"
    )
    return result


def _rows(workbook, sheet_name: str):
    return list(workbook[sheet_name].iter_rows(min_row=2, values_only=True))


SCALAR_AND_OUT_SOURCE = """
import json
from pathlib import Path

import torch
from openpyxl import load_workbook

from aten_size_profiler import AtenSizeMode, export_collector

def _rows(workbook, sheet_name):
    return list(workbook[sheet_name].iter_rows(min_row=2, values_only=True))

output_path = Path(r"{output_path}")
left = torch.ones(2)
right = torch.ones(2)
with AtenSizeMode(record_scalar_args=True) as mode:
    torch.add(left, right, alpha=2)
    torch.add(left, right, alpha=3)
    output = torch.empty(2)
    torch.add(left, right, alpha=4, out=output)

add_calls = [
    call for call in mode.calls if call.op_name == "aten.add.Tensor"
]
out_calls = [
    call for call in mode.calls if call.op_name == "aten.add.out"
]
assert len(add_calls) == 2
assert {{call.scalar_args[0][1] for call in add_calls}} == {{json.dumps(2), json.dumps(3)}}
assert len(out_calls) == 1
assert len(out_calls[0].inputs) == 2
assert all(tensor.path != "kwargs['out']" for tensor in out_calls[0].inputs)
assert sum(tensor.logical_bytes for tensor in out_calls[0].inputs) == 16

export_collector(mode.collector, output_path)
workbook = load_workbook(output_path, read_only=True, data_only=True)
assert workbook.sheetnames == {sheets!r}
unique_rows = [
    row for row in _rows(workbook, "unique_shapes")
    if row[0] == "aten.add.Tensor"
]
assert len(unique_rows) == 2
assert {{json.loads(row[8])[0][1] for row in unique_rows}} == {{"2", "3"}}
detail_rows = [
    row for row in _rows(workbook, "detail")
    if row[1] == "aten.add.out"
]
assert detail_rows[0][2] == 2
assert json.loads(detail_rows[0][14])[0][1] == "4"
workbook.close()

with AtenSizeMode(record_scalar_args=False) as no_scalar_mode:
    torch.add(left, right, alpha=2)
    torch.add(left, right, alpha=3)
assert len([
    call for call in no_scalar_mode.calls
    if call.op_name == "aten.add.Tensor"
]) == 2
"""


LONG_CELL_SOURCE = """
from pathlib import Path

from openpyxl import load_workbook

from aten_size_profiler import AtenCollector, TensorInfo, export_collector

output_path = Path(r"{output_path}")
collector = AtenCollector()
inputs = tuple(
    TensorInfo(
        path=f"args[0][{{index}}]",
        shape=(1,),
        dtype="torch.float32",
        device="cpu",
        stride=(1,),
        numel=1,
        element_size=4,
        logical_bytes=4,
    )
    for index in range(100)
)
collector.record("aten.test.default", inputs)
export_collector(collector, output_path)
workbook = load_workbook(output_path, read_only=True, data_only=True)
detail_row = next(workbook["detail"].iter_rows(min_row=2, values_only=True))
assert len(__import__("json").loads(detail_row[3])) == 100
input_rows = list(workbook["input_tensors"].iter_rows(min_row=2, values_only=True))
assert len(input_rows) == 100
assert len(input_rows) == detail_row[2]
workbook.close()

too_long = AtenCollector()
too_long.record(
    "aten.test.default",
    tuple(
        TensorInfo(
            path=f"args[0][{{index}}]-" + "x" * 8,
            shape=(1,),
            dtype="torch.float32",
            device="cpu",
            stride=(1,),
            numel=1,
            element_size=4,
            logical_bytes=4,
        )
        for index in range(4000)
    ),
)
too_long_path = export_collector(too_long, output_path.with_name("too_long.xlsx"))
too_long_workbook = load_workbook(too_long_path, read_only=True, data_only=True)
too_long_detail_row = next(
    too_long_workbook["detail"].iter_rows(min_row=2, values_only=True)
)
assert too_long_detail_row[3].endswith("...<TRUNCATED>")
assert len(too_long_detail_row[3]) == 32767
too_long_metadata = {{
    row[0]: row[1]
    for row in too_long_workbook["metadata"].iter_rows(min_row=2, values_only=True)
}}
assert too_long_metadata["truncated_cells"] > 0
assert "detail" in __import__("json").loads(too_long_metadata["truncated_sheets"])
too_long_workbook.close()
"""


MODE_STACK_SOURCE = """
import torch
from torch.utils._python_dispatch import TorchDispatchMode

import aten_size_profiler as profiler

class UserMode(TorchDispatchMode):
    def __init__(self):
        super().__init__()
        self.calls = []

    def __torch_dispatch__(self, func, types, args=(), kwargs=None):
        self.calls.append(str(func))
        return func(*args, **(kwargs or {}))

user_mode = UserMode()
user_mode.__enter__()
torch.ones(2)
before = len(user_mode.calls)
assert profiler.stop_profiler(export=False) is None
torch.ones(2)
assert len(user_mode.calls) == before + 1
user_mode.__exit__(None, None, None)
"""


SPAWN_SOURCE = """
import os
from pathlib import Path

import torch
import torch.distributed as dist
import torch.multiprocessing as mp
from openpyxl import load_workbook
from torch import nn

OUTPUT_DIR = Path(r"{output_dir}")
INIT_FILE = r"{init_file}"

def rows(workbook, sheet_name):
    return list(workbook[sheet_name].iter_rows(min_row=2, values_only=True))

def worker(rank):
    import aten_size_profiler as profiler

    dist.init_process_group(
        "gloo",
        init_method=f"file://{{INIT_FILE}}",
        rank=rank,
        world_size=2,
    )

    class Model(nn.Module):
        def __init__(self):
            super().__init__()
            self.weight = nn.Parameter(torch.ones(4, 4))

        def forward(self, value):
            return torch.mm(value, self.weight)

    model = Model()
    value = torch.ones(2 + rank * 3, 4)
    model(value)
    dist.destroy_process_group()

    report = profiler.stop_profiler()
    assert report.name == f"spawn_rank{{rank}}.xlsx"
    workbook = load_workbook(report, read_only=True, data_only=True)
    assert workbook.sheetnames == {sheets!r}
    detail_rows = rows(workbook, "detail")
    input_rows = rows(workbook, "input_tensors")
    metadata = dict(rows(workbook, "metadata"))
    assert detail_rows
    assert {{row[13] for row in detail_rows}} == {{rank}}
    assert {{row[12] for row in input_rows}} == {{rank}}
    assert {{row[12] for row in detail_rows}} == {{0}}
    assert metadata["rank"] == rank
    assert metadata["global_rank"] == rank
    assert metadata["world_size"] == 2
    assert metadata["distributed_initialized"] is True
    expected_shape = f"[{{2 + rank * 3}},4]"
    other_shape = f"[{{2 + (1 - rank) * 3}},4]"
    shapes = {{row[4] for row in input_rows}}
    assert expected_shape in shapes
    assert other_shape not in shapes
    workbook.close()

if __name__ == "__main__":
    mp.spawn(worker, nprocs=2, join=True)
"""


AUTO_EXIT_SOURCE = """
from pathlib import Path

import torch
import torch.distributed as dist

import aten_size_profiler

dist.init_process_group(
    "gloo",
    init_method="file://{init_file}",
    rank=0,
    world_size=1,
)
torch.ones(2)
dist.destroy_process_group()
"""


def run() -> None:
    with tempfile.TemporaryDirectory(prefix="aten-cross-review-batch4-") as path:
        root = Path(path)
        scalar_report = root / "scalar.xlsx"
        _run(
            "scalar-and-out",
            SCALAR_AND_OUT_SOURCE.format(
                output_path=scalar_report,
                sheets=SHEETS,
            ),
            ATEN_SIZE_PROFILER_ENABLED="false",
        )

        long_report = root / "long.xlsx"
        _run(
            "long-cell",
            LONG_CELL_SOURCE.format(output_path=long_report),
            ATEN_SIZE_PROFILER_ENABLED="false",
        )

        _run(
            "mode-stack",
            MODE_STACK_SOURCE,
            ATEN_SIZE_PROFILER_ENABLED="true",
        )

        init_file = root / "spawn-init"
        spawn_output = root / "spawn-reports"
        spawn_output.mkdir()
        _run_file(
            "mp-spawn-distributed",
            SPAWN_SOURCE.format(
                output_dir=spawn_output,
                init_file=init_file,
                sheets=SHEETS,
            ),
            root,
            ATEN_SIZE_PROFILER_ENABLED="true",
            ATEN_SIZE_PROFILER_OUTPUT_DIR=str(spawn_output),
            ATEN_SIZE_PROFILER_OUTPUT_FILENAME="spawn.xlsx",
            ATEN_SIZE_PROFILER_TARGET_ITER="0",
            ATEN_SIZE_PROFILER_ITERATION_MODE="forward",
        )
        assert (spawn_output / "spawn_rank0.xlsx").is_file()
        assert (spawn_output / "spawn_rank1.xlsx").is_file()

        auto_output = root / "auto.xlsx"
        auto_init = root / "auto-init"
        _run(
            "distributed-auto-exit",
            AUTO_EXIT_SOURCE.format(init_file=auto_init),
            ATEN_SIZE_PROFILER_ENABLED="true",
            ATEN_SIZE_PROFILER_OUTPUT_DIR=str(root),
            ATEN_SIZE_PROFILER_OUTPUT_FILENAME=auto_output.name,
        )
        workbook = load_workbook(auto_output, read_only=True, data_only=True)
        metadata = dict(_rows(workbook, "metadata"))
        assert metadata["distributed_initialized"] is True
        workbook.close()

    print("CROSS REVIEW BATCH 4 TEST PASS")


if __name__ == "__main__":
    run()
