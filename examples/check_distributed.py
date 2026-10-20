from __future__ import annotations

import os
from pathlib import Path

import torch
import torch.distributed as dist
from openpyxl import load_workbook
from torch import nn


OUTPUT_DIR = Path(
    os.environ.get(
        "ATEN_SIZE_PROFILER_OUTPUT_DIR",
        "./aten_size_reports/stage5_validation/latest",
    )
)
os.environ["ATEN_SIZE_PROFILER_ENABLED"] = "true"
os.environ["ATEN_SIZE_PROFILER_OUTPUT_DIR"] = str(OUTPUT_DIR)
os.environ["ATEN_SIZE_PROFILER_OUTPUT_FILENAME"] = "aten_size_profile_size.xlsx"
os.environ["ATEN_SIZE_PROFILER_TARGET_ITER"] = "1"
os.environ["ATEN_SIZE_PROFILER_ITERATION_MODE"] = "forward"


def _rows(workbook, sheet_name: str):
    return list(workbook[sheet_name].iter_rows(min_row=2, values_only=True))


def run() -> None:
    dist.init_process_group("gloo")
    rank = dist.get_rank()
    world_size = dist.get_world_size()
    assert world_size >= 2

    class Model(nn.Module):
        def __init__(self):
            super().__init__()
            self.weight = nn.Parameter(torch.randn(4, 4))

        def forward(self, inputs):
            return torch.mm(inputs, self.weight)

    torch.manual_seed(2026)
    model = Model()
    rows_by_iteration = (2 + rank * 3, 3 + rank * 3, 4 + rank * 3)
    inputs = [torch.randn(rows, 4) for rows in rows_by_iteration]
    baseline = [model(value).detach().clone() for value in inputs]

    import aten_size_profiler as profiler

    rank_info = profiler.get_rank_info()
    assert rank_info.rank == rank
    assert rank_info.world_size == world_size
    assert rank_info.initialized

    for value, expected in zip(inputs, baseline):
        actual = model(value)
        torch.testing.assert_close(actual, expected)

    output_path = profiler.stop_profiler()
    assert output_path is not None
    assert output_path.name.endswith(f"_rank{rank}.xlsx")
    assert output_path.parent == OUTPUT_DIR

    workbook = load_workbook(output_path, read_only=True, data_only=True)
    assert workbook.sheetnames == [
        "summary",
        "unique_shapes",
        "detail",
        "input_tensors",
        "metadata",
    ]
    detail_rows = _rows(workbook, "detail")
    input_rows = _rows(workbook, "input_tensors")
    metadata = {
        row[0]: row[1]
        for row in _rows(workbook, "metadata")
    }
    assert detail_rows
    assert all(row[12] == 1 for row in detail_rows)
    assert all(row[11] == 1 for row in input_rows)
    assert all(row[13] == rank for row in detail_rows)
    assert all(row[12] == rank for row in input_rows)
    assert metadata["rank"] == rank
    assert metadata["global_rank"] == rank
    assert metadata["world_size"] == world_size
    assert metadata["distributed_initialized"] is True
    target_shape = f"[{rows_by_iteration[1]},4]"
    other_rank_target_shape = f"[{rows_by_iteration[1] + (3 if rank == 0 else -3)},4]"
    input_shapes = {row[4] for row in input_rows}
    assert target_shape in input_shapes
    assert other_rank_target_shape not in input_shapes
    workbook.close()

    print(f"RANK {rank} PASS: {output_path}")
    dist.destroy_process_group()


if __name__ == "__main__":
    run()
