from __future__ import annotations

import os
import subprocess
import sys
import tempfile
import textwrap
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]


def _environment(**updates: str) -> dict[str, str]:
    environment = {
        key: value
        for key, value in os.environ.items()
        if not key.startswith("ATEN_SIZE_PROFILER_")
    }
    environment["PYTHONPATH"] = str(ROOT)
    environment.update(updates)
    return environment


def _run(name: str, source: str, **updates: str) -> subprocess.CompletedProcess[str]:
    result = subprocess.run(
        [sys.executable, "-c", textwrap.dedent(source)],
        cwd=ROOT,
        env=_environment(**updates),
        text=True,
        capture_output=True,
        timeout=240,
    )
    assert result.returncode == 0, (
        f"{name} failed with code {result.returncode}\n"
        f"stdout:\n{result.stdout}\nstderr:\n{result.stderr}"
    )
    return result


MULTI_FORWARD_SOURCE = """
import os

import torch
from openpyxl import load_workbook
from torch import nn

import aten_size_profiler as profiler

torch.manual_seed(2026)
kind = os.environ["CHECK_KIND"]
model = nn.Linear(4, 4)
optimizer = torch.optim.SGD(model.parameters(), lr=0.1)
batch_sizes = iter(range(2, 100))

for _ in range(3):
    if kind == "accumulation":
        for _ in range(2):
            model(torch.randn(next(batch_sizes), 4)).sum().backward()
    else:
        first = model(torch.randn(next(batch_sizes), 4))
        second = model(torch.randn(next(batch_sizes), 4))
        (first.sum() + second.sum()).backward()
    optimizer.step()
    optimizer.zero_grad()

recorded = [
    call.inputs[1].shape[0]
    for call in profiler.get_collector().calls
    if call.op_name == "aten.addmm.default"
]
assert recorded == [4, 5], recorded
assert profiler.get_iteration_status()["phase"] == "FINISHED"

output_path = profiler.stop_profiler()
workbook = load_workbook(output_path, read_only=True, data_only=True)
metadata = {
    row[0]: row[1]
    for row in workbook["metadata"].iter_rows(min_row=2, values_only=True)
}
assert metadata["iteration_complete"] is True
iterations = {
    row[12] for row in workbook["detail"].iter_rows(min_row=2, values_only=True)
}
assert iterations == {1}, iterations
workbook.close()
"""


LBFGS_SOURCE = """
import torch
from torch import nn

import aten_size_profiler as profiler

torch.manual_seed(2026)
model = nn.Linear(4, 1)
optimizer = torch.optim.LBFGS(model.parameters(), max_iter=4)
inputs = torch.randn(8, 4)
targets = torch.randn(8, 1)
evaluations = 0

def closure():
    global evaluations
    evaluations += 1
    optimizer.zero_grad()
    loss = ((model(inputs) - targets) ** 2).mean()
    loss.backward()
    return loss

evaluations_per_step = []
for _ in range(3):
    before = evaluations
    optimizer.step(closure)
    evaluations_per_step.append(evaluations - before)

recorded = [
    call
    for call in profiler.get_collector().calls
    if call.op_name == "aten.addmm.default"
]
assert evaluations_per_step[1] > 1, evaluations_per_step
assert len(recorded) == evaluations_per_step[1], (
    len(recorded),
    evaluations_per_step,
)
assert profiler.get_iteration_status()["phase"] == "FINISHED"
profiler.stop_profiler(export=False)
"""


FOREACH_LONG_CELL_SOURCE = """
import json

import torch
from openpyxl import load_workbook
from torch import nn

import aten_size_profiler as profiler

model = nn.Sequential(*[nn.Linear(8, 8) for _ in range(400)])
optimizer = torch.optim.Adam(model.parameters(), lr=1e-3, foreach=True)
for _ in range(3):
    model(torch.randn(2, 8)).sum().backward()
    optimizer.step()
    optimizer.zero_grad()

output_path = profiler.stop_profiler()
assert profiler.get_status()["export_count"] == 1
workbook = load_workbook(output_path, read_only=True, data_only=True)
metadata = {
    row[0]: row[1]
    for row in workbook["metadata"].iter_rows(min_row=2, values_only=True)
}
assert metadata["truncated_cells"] > 0
assert "unique_shapes" in json.loads(metadata["truncated_sheets"])
truncated = [
    value
    for row in workbook["unique_shapes"].iter_rows(min_row=2, values_only=True)
    for value in row
    if isinstance(value, str) and value.endswith("...<TRUNCATED>")
]
assert truncated
assert all(len(value) == 32767 for value in truncated)
foreach_rows = [
    row
    for row in workbook["summary"].iter_rows(min_row=2, values_only=True)
    if row[0].startswith("aten._foreach_")
]
assert foreach_rows
workbook.close()
"""


EXPORT_RETRY_SOURCE = """
import torch

import aten_size_profiler as profiler

torch.ones(2) + 1
errors = []
for _ in range(2):
    try:
        profiler.stop_profiler()
    except UnboundLocalError:
        raise
    except OSError as error:
        errors.append(type(error).__name__)
assert len(errors) == 2, errors
"""


COND_SOURCE = """
import torch

import aten_size_profiler as profiler

values = torch.randn(3)
result = torch.cond(
    values.sum() > 0,
    lambda tensor: tensor.sin(),
    lambda tensor: tensor.cos(),
    (values,),
)
expected = values.sin() if values.sum() > 0 else values.cos()
torch.testing.assert_close(result, expected)
recorded = {call.op_name for call in profiler.get_collector().calls}
assert recorded & {"aten.sin.default", "aten.cos.default"}, recorded
profiler.stop_profiler(export=False)
"""


COND_WAITING_SOURCE = """
import torch

import aten_size_profiler as profiler

assert profiler.get_iteration_status()["phase"] == "WAITING"
assert not profiler.is_recording()
values = torch.randn(3)
result = torch.cond(
    values.sum() > 0,
    lambda tensor: tensor.sin(),
    lambda tensor: tensor.cos(),
    (values,),
)
expected = values.sin() if values.sum() > 0 else values.cos()
torch.testing.assert_close(result, expected)
assert profiler.get_collector().calls == []
profiler.stop_profiler(export=False)
"""


FLEX_ATTENTION_SOURCE = """
import warnings

import torch

import aten_size_profiler as profiler

warnings.filterwarnings("ignore", category=UserWarning)
from torch.nn.attention.flex_attention import flex_attention

query = torch.randn(1, 2, 16, 8)
result = flex_attention(query, query, query)
assert tuple(result.shape) == (1, 2, 16, 8)
assert profiler.get_collector().calls
profiler.stop_profiler(export=False)
"""


def run() -> None:
    with tempfile.TemporaryDirectory(prefix="aten-cross-review-round2-") as path:
        root = Path(path)
        for kind in ("accumulation", "siamese"):
            _run(
                f"multi-forward-{kind}",
                MULTI_FORWARD_SOURCE,
                CHECK_KIND=kind,
                ATEN_SIZE_PROFILER_ENABLED="true",
                ATEN_SIZE_PROFILER_ITERATION_MODE="optimizer_step",
                ATEN_SIZE_PROFILER_TARGET_ITER="1",
                ATEN_SIZE_PROFILER_OUTPUT_DIR=str(root),
                ATEN_SIZE_PROFILER_OUTPUT_FILENAME=f"{kind}.xlsx",
            )

        _run(
            "lbfgs-closure",
            LBFGS_SOURCE,
            ATEN_SIZE_PROFILER_ENABLED="true",
            ATEN_SIZE_PROFILER_ITERATION_MODE="optimizer_step",
            ATEN_SIZE_PROFILER_TARGET_ITER="1",
        )

        _run(
            "foreach-long-cell",
            FOREACH_LONG_CELL_SOURCE,
            ATEN_SIZE_PROFILER_ENABLED="true",
            ATEN_SIZE_PROFILER_ITERATION_MODE="optimizer_step",
            ATEN_SIZE_PROFILER_TARGET_ITER="1",
            ATEN_SIZE_PROFILER_OUTPUT_DIR=str(root),
            ATEN_SIZE_PROFILER_OUTPUT_FILENAME="foreach.xlsx",
        )

        blocked_output_dir = root / "not-a-directory"
        blocked_output_dir.write_text("", encoding="utf-8")
        _run(
            "export-retry",
            EXPORT_RETRY_SOURCE,
            ATEN_SIZE_PROFILER_ENABLED="true",
            ATEN_SIZE_PROFILER_OUTPUT_DIR=str(blocked_output_dir),
        )

        _run(
            "cond-recording",
            COND_SOURCE,
            ATEN_SIZE_PROFILER_ENABLED="true",
        )
        _run(
            "cond-waiting",
            COND_WAITING_SOURCE,
            ATEN_SIZE_PROFILER_ENABLED="true",
            ATEN_SIZE_PROFILER_ITERATION_MODE="forward",
            ATEN_SIZE_PROFILER_TARGET_ITER="5",
        )
        _run(
            "flex-attention",
            FLEX_ATTENTION_SOURCE,
            ATEN_SIZE_PROFILER_ENABLED="true",
        )

    print("CROSS REVIEW ROUND 2 TEST PASS")


if __name__ == "__main__":
    run()
