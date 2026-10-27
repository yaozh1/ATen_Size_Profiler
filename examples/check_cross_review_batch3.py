from __future__ import annotations

import os
import subprocess
import sys
import tempfile
import textwrap
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SHEETS = [
    "summary",
    "unique_shapes",
    "detail",
    "input_tensors",
    "metadata",
]


def _environment(
    output_dir: Path,
    output_filename: str,
    target_iter: int,
    iteration_mode: str,
    **updates: str,
) -> dict[str, str]:
    environment = {
        key: value
        for key, value in os.environ.items()
        if not key.startswith("ATEN_SIZE_PROFILER_")
    }
    environment.update(
        {
            "PYTHONPATH": str(ROOT),
            "ATEN_SIZE_PROFILER_ENABLED": "true",
            "ATEN_SIZE_PROFILER_OUTPUT_DIR": str(output_dir),
            "ATEN_SIZE_PROFILER_OUTPUT_FILENAME": output_filename,
            "ATEN_SIZE_PROFILER_TARGET_ITER": str(target_iter),
            "ATEN_SIZE_PROFILER_ITERATION_MODE": iteration_mode,
        }
    )
    environment.update(updates)
    return environment


def _run(
    name: str,
    source: str,
    output_dir: Path,
    output_filename: str,
    target_iter: int,
    iteration_mode: str,
    **updates: str,
) -> Path:
    result = subprocess.run(
        [sys.executable, "-c", textwrap.dedent(source)],
        cwd=ROOT,
        env=_environment(
            output_dir,
            output_filename,
            target_iter,
            iteration_mode,
            **updates,
        ),
        text=True,
        capture_output=True,
        timeout=180,
    )
    assert result.returncode == 0, (
        f"{name} failed with code {result.returncode}\n"
        f"stdout:\n{result.stdout}\nstderr:\n{result.stderr}"
    )
    path = output_dir / output_filename
    assert path.is_file(), f"{name} report not found: {path}"
    return path


def _rows(workbook, sheet_name: str):
    return list(workbook[sheet_name].iter_rows(min_row=2, values_only=True))


def _validate_report(
    path: Path,
    *,
    expected_iteration: int,
    complete: bool,
    has_detail: bool,
    excluded_operator: str | None = None,
) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    assert workbook.sheetnames == SHEETS
    detail_rows = _rows(workbook, "detail")
    metadata = dict(_rows(workbook, "metadata"))
    assert metadata["iteration_complete"] is complete
    if has_detail:
        assert detail_rows
        assert {row[12] for row in detail_rows} == {expected_iteration}
        assert metadata["iteration_phase"] == "FINISHED"
    else:
        assert not detail_rows
        assert metadata["iteration_phase"] != "FINISHED"
    if excluded_operator is not None:
        assert all(row[1] != excluded_operator for row in detail_rows)
    workbook.close()


FORWARD_SOURCE = """
import torch
from torch import nn

import aten_size_profiler as profiler

model = nn.Linear(4, 4)
criterion = nn.MSELoss()
for rows in (2, 3, 5):
    value = torch.randn(rows, 4)
    output = model(value)
    criterion(output, torch.zeros_like(output))
assert profiler.get_iteration_status()["current_iter"] == 2
assert profiler.get_iteration_status()["phase"] == "FINISHED"
profiler.stop_profiler()
"""


AMP_SOURCE = """
import os
import torch
from torch import nn

import aten_size_profiler as profiler

target = int(os.environ["ATEN_SIZE_PROFILER_TARGET_ITER"])
skip_count = int(os.environ["SKIP_COUNT"])
steps = skip_count + 1
model = nn.Linear(4, 4)
optimizer = torch.optim.SGD(model.parameters(), lr=0.1)
scaler = torch.amp.GradScaler("cpu", enabled=True)
for index in range(steps):
    optimizer.zero_grad(set_to_none=True)
    output = model(torch.randn(2, 4))
    loss = output.square().mean()
    scaler.scale(loss).backward()
    if index < skip_count:
        for parameter in model.parameters():
            parameter.grad.fill_(float("inf"))
    scaler.step(optimizer)
    scaler.update()
status = profiler.get_iteration_status()
assert status["current_iter"] == steps
assert status["phase"] == "FINISHED"
if target < steps and target >= skip_count:
    assert profiler.get_collector().calls
    assert all(
        call.iteration == target
        for call in profiler.get_collector().calls
    )
profiler.stop_profiler()
"""


EXCEPTION_SOURCE = """
import os
import torch
from torch import nn

import aten_size_profiler as profiler

stage = os.environ["FAIL_STAGE"]

class FailingSGD(torch.optim.SGD):
    def step(self, *args, **kwargs):
        raise RuntimeError("expected optimizer failure")

model = nn.Linear(4, 4)
optimizer_type = FailingSGD if stage == "optimizer" else torch.optim.SGD
optimizer = optimizer_type(model.parameters(), lr=0.1)
try:
    output = model(torch.randn(2, 4))
    if stage == "loss":
        raise RuntimeError("expected loss failure")
    loss = output.detach().sum() if stage == "backward" else output.square().mean()
    loss.backward()
    optimizer.step()
except RuntimeError:
    pass
assert profiler.get_iteration_status()["phase"] in {"RECORDING", "WAITING"}
profiler.stop_profiler()
"""


RETRY_SOURCE = """
import torch
from torch import nn

import aten_size_profiler as profiler

model = nn.Linear(4, 4)
optimizer = torch.optim.SGD(model.parameters(), lr=0.1)
try:
    model(torch.randn(2, 4))
    raise RuntimeError("expected loss failure")
except RuntimeError:
    pass
optimizer.zero_grad(set_to_none=True)
output = model(torch.randn(3, 4))
loss = output.square().mean()
loss.backward()
optimizer.step()
assert profiler.get_iteration_status()["phase"] == "FINISHED"
profiler.stop_profiler()
"""


def run() -> None:
    with tempfile.TemporaryDirectory(prefix="aten-cross-review-batch3-") as path:
        output_dir = Path(path)
        forward_report = _run(
            "forward-root-module",
            FORWARD_SOURCE,
            output_dir,
            "forward.xlsx",
            1,
            "forward",
        )
        _validate_report(
            forward_report,
            expected_iteration=1,
            complete=True,
            has_detail=True,
            excluded_operator="aten.mse_loss.default",
        )

        for name, target, skip_count, expected_detail in (
            ("amp_normal", 0, 0, True),
            ("amp_consecutive_skip", 2, 2, True),
            ("amp_target_skip", 1, 1, True),
        ):
            report = _run(
                name,
                AMP_SOURCE,
                output_dir,
                f"{name}.xlsx",
                target,
                "optimizer_step",
                SKIP_COUNT=str(skip_count),
            )
            _validate_report(
                report,
                expected_iteration=target,
                complete=True,
                has_detail=expected_detail,
            )

        for stage in ("loss", "backward", "optimizer"):
            report = _run(
                f"exception-{stage}",
                EXCEPTION_SOURCE,
                output_dir,
                f"exception_{stage}.xlsx",
                0,
                "optimizer_step",
                FAIL_STAGE=stage,
            )
            _validate_report(
                report,
                expected_iteration=0,
                complete=False,
                has_detail=False,
            )

        retry_report = _run(
            "exception-retry",
            RETRY_SOURCE,
            output_dir,
            "retry.xlsx",
            0,
            "optimizer_step",
        )
        _validate_report(
            retry_report,
            expected_iteration=0,
            complete=True,
            has_detail=True,
        )
    print("CROSS REVIEW BATCH 3 TEST PASS")


if __name__ == "__main__":
    run()
