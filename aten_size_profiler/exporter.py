from __future__ import annotations

import json
import platform
import sys
from collections import defaultdict
from collections.abc import Mapping
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import torch
from openpyxl import Workbook

from . import __version__
from .collector import AtenCollector, CallRecord, TensorInfo
from .config import validate_output_path


TOOL_VERSION = __version__
MAX_EXCEL_ROWS = 1_048_576
MAX_EXCEL_CELL_LENGTH = 32_767


@dataclass
class _ExportStats:
    truncated_cells: int = 0
    truncated_sheets: set[str] = field(default_factory=set)


def _format_sequence(values: tuple[Any, ...] | list[Any] | None) -> str:
    if values is None:
        return ""
    return json.dumps(list(values), ensure_ascii=False, separators=(",", ":"))


def _format_shape(shape: tuple[int, ...]) -> str:
    return _format_sequence(shape)


def _format_stride(stride: tuple[int, ...] | None) -> str:
    return _format_sequence(stride)


def _format_tensor_list(
    inputs: tuple[TensorInfo, ...],
    attribute: str,
) -> str:
    return _format_sequence([getattr(tensor, attribute) for tensor in inputs])


def _format_scalar_args(
    scalar_args: tuple[tuple[str, str], ...],
) -> str:
    return _format_sequence(
        [[path, value] for path, value in scalar_args]
    )


TRUNCATION_MARKER = "...<TRUNCATED>"


def _append_row(
    worksheet: Any,
    row: list[Any],
    sheet_name: str,
    row_number: int,
    stats: _ExportStats,
) -> None:
    safe_row = list(row)
    for column_number, value in enumerate(row, start=1):
        if isinstance(value, str) and len(value) > MAX_EXCEL_CELL_LENGTH:
            truncated_length = MAX_EXCEL_CELL_LENGTH - len(TRUNCATION_MARKER)
            safe_row[column_number - 1] = value[:truncated_length] + TRUNCATION_MARKER
            stats.truncated_cells += 1
            stats.truncated_sheets.add(sheet_name)
    worksheet.append(safe_row)


def _format_metadata_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, Path):
        return str(value)
    return json.dumps(value, ensure_ascii=False, default=str)


def _default_metadata(
    collector: AtenCollector,
    output_path: Path,
) -> dict[str, Any]:
    devices = sorted(
        {
            tensor.device
            for call in collector.calls
            for tensor in call.inputs
        }
    )
    return {
        "tool_version": TOOL_VERSION,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "output_file": str(output_path),
        "python_version": sys.version.split()[0],
        "platform": platform.platform(),
        "pytorch_version": torch.__version__,
        "cuda_version": torch.version.cuda or "",
        "hip_version": torch.version.hip or "",
        "cuda_available": torch.cuda.is_available(),
        "devices": devices,
        "call_count": len(collector.calls),
        "iteration": "not configured",
        "rank": "not configured",
        "record_stride": True,
        "record_stack": False,
        "record_scalar_args": True,
    }


def _write_summary_sheet(
    workbook: Workbook,
    calls: list[CallRecord],
    stats: _ExportStats,
) -> None:
    worksheet = workbook.create_sheet("summary")
    worksheet.freeze_panes = "A2"
    _append_row(
        worksheet,
        [
            "op_name",
            "call_count",
            "unique_shape_count",
            "input_tensor_count_total",
            "input_logical_bytes_total",
        ],
        "summary",
        1,
        stats,
    )
    row_number = 2

    grouped_calls: dict[str, list[CallRecord]] = defaultdict(list)
    for call in calls:
        grouped_calls[call.op_name].append(call)

    for op_name in sorted(grouped_calls):
        operator_calls = grouped_calls[op_name]
        unique_keys = {_unique_shape_key(call) for call in operator_calls}
        input_tensors = [
            tensor
            for call in operator_calls
            for tensor in call.inputs
        ]
        _append_row(
            worksheet,
            [
                op_name,
                len(operator_calls),
                len(unique_keys),
                len(input_tensors),
                sum(tensor.logical_bytes for tensor in input_tensors),
            ],
            "summary",
            row_number,
            stats,
        )
        row_number += 1


def _unique_shape_key(call: CallRecord) -> tuple[Any, ...]:
    return tuple(
        (
            tensor.shape,
            tensor.dtype,
            tensor.device,
            tensor.stride,
        )
        for tensor in call.inputs
    ) + (call.scalar_args,)


def _write_unique_shapes_sheet(
    workbook: Workbook,
    calls: list[CallRecord],
    stats: _ExportStats,
) -> None:
    worksheet = workbook.create_sheet("unique_shapes")
    worksheet.freeze_panes = "A2"
    _append_row(
        worksheet,
        [
            "op_name",
            "input_shapes",
            "input_dtypes",
            "input_devices",
            "input_strides",
            "count",
            "op_ratio",
            "overall_ratio",
            "scalar_args",
        ],
        "unique_shapes",
        1,
        stats,
    )
    row_number = 2

    grouped_calls: dict[str, list[CallRecord]] = defaultdict(list)
    for call in calls:
        grouped_calls[call.op_name].append(call)
    total_calls = len(calls)

    for op_name in sorted(grouped_calls):
        operator_calls = grouped_calls[op_name]
        shape_groups: dict[tuple[Any, ...], CallRecord] = {}
        shape_counts: dict[tuple[Any, ...], int] = defaultdict(int)
        for call in operator_calls:
            key = _unique_shape_key(call)
            shape_groups.setdefault(key, call)
            shape_counts[key] += 1

        for key in sorted(shape_counts, key=repr):
            representative = shape_groups[key]
            _append_row(
                worksheet,
                [
                    op_name,
                    _format_tensor_list(representative.inputs, "shape"),
                    _format_tensor_list(representative.inputs, "dtype"),
                    _format_tensor_list(representative.inputs, "device"),
                    _format_tensor_list(representative.inputs, "stride"),
                    shape_counts[key],
                    shape_counts[key] / len(operator_calls),
                    shape_counts[key] / total_calls if total_calls else 0,
                    _format_scalar_args(representative.scalar_args),
                ],
                "unique_shapes",
                row_number,
                stats,
            )
            row_number += 1


def _write_detail_sheet(
    workbook: Workbook,
    calls: list[CallRecord],
    iteration: Any,
    rank: Any,
    stats: _ExportStats,
) -> None:
    worksheet = workbook.create_sheet("detail")
    worksheet.freeze_panes = "A2"
    _append_row(
        worksheet,
        [
            "call_id",
            "op_name",
            "input_count",
            "input_paths",
            "input_shapes",
            "input_dtypes",
            "input_devices",
            "input_strides",
            "input_numels",
            "input_element_sizes",
            "input_logical_bytes",
            "stack",
            "iteration",
            "rank",
            "scalar_args",
        ],
        "detail",
        1,
        stats,
    )
    row_number = 2

    for call in calls:
        call_iteration = (
            call.iteration if call.iteration is not None else iteration
        )
        _append_row(
            worksheet,
            [
                call.call_id,
                call.op_name,
                len(call.inputs),
                _format_tensor_list(call.inputs, "path"),
                _format_tensor_list(call.inputs, "shape"),
                _format_tensor_list(call.inputs, "dtype"),
                _format_tensor_list(call.inputs, "device"),
                _format_tensor_list(call.inputs, "stride"),
                _format_tensor_list(call.inputs, "numel"),
                _format_tensor_list(call.inputs, "element_size"),
                sum(tensor.logical_bytes for tensor in call.inputs),
                "".join(call.stack),
                call_iteration,
                rank,
                _format_scalar_args(call.scalar_args),
            ],
            "detail",
            row_number,
            stats,
        )
        row_number += 1


def _write_input_tensors_sheet(
    workbook: Workbook,
    calls: list[CallRecord],
    iteration: Any,
    rank: Any,
    stats: _ExportStats,
) -> None:
    worksheet = workbook.create_sheet("input_tensors")
    worksheet.freeze_panes = "A2"
    _append_row(
        worksheet,
        [
            "call_id",
            "op_name",
            "input_index",
            "input_path",
            "shape",
            "dtype",
            "device",
            "stride",
            "numel",
            "element_size",
            "logical_bytes",
            "iteration",
            "rank",
        ],
        "input_tensors",
        1,
        stats,
    )
    row_number = 2

    for call in calls:
        call_iteration = (
            call.iteration if call.iteration is not None else iteration
        )
        for input_index, tensor in enumerate(call.inputs):
            _append_row(
                worksheet,
                [
                    call.call_id,
                    call.op_name,
                    input_index,
                    tensor.path,
                    _format_shape(tensor.shape),
                    tensor.dtype,
                    tensor.device,
                    _format_stride(tensor.stride),
                    tensor.numel,
                    tensor.element_size,
                    tensor.logical_bytes,
                    call_iteration,
                    rank,
                ],
                "input_tensors",
                row_number,
                stats,
            )
            row_number += 1


def _write_metadata_sheet(
    workbook: Workbook,
    metadata: Mapping[str, Any],
    stats: _ExportStats,
) -> None:
    worksheet = workbook.create_sheet("metadata")
    worksheet.freeze_panes = "A2"
    _append_row(
        worksheet,
        ["key", "value"],
        "metadata",
        1,
        stats,
    )
    row_number = 2
    for key in sorted(metadata):
        _append_row(
            worksheet,
            [key, _format_metadata_value(metadata[key])],
            "metadata",
            row_number,
            stats,
        )
        row_number += 1


def _validate_row_limits(row_counts: Mapping[str, int]) -> None:
    for sheet_name, row_count in row_counts.items():
        if row_count > MAX_EXCEL_ROWS:
            raise ValueError(
                f"Excel sheet {sheet_name!r} would contain {row_count} rows, "
                f"exceeding the limit of {MAX_EXCEL_ROWS}"
            )


def export_collector(
    collector: AtenCollector,
    output_path: str | Path,
    *,
    metadata: Mapping[str, Any] | None = None,
    iteration: Any = None,
    rank: Any = None,
) -> Path:
    destination = validate_output_path(output_path)
    destination.parent.mkdir(parents=True, exist_ok=True)

    report_metadata = _default_metadata(collector, destination)
    if metadata is not None:
        report_metadata.update(metadata)
    if iteration is not None:
        report_metadata["iteration"] = iteration
    if rank is not None:
        report_metadata["rank"] = rank

    grouped_calls: dict[str, list[CallRecord]] = defaultdict(list)
    for call in collector.calls:
        grouped_calls[call.op_name].append(call)
    unique_shape_count = 0
    for operator_calls in grouped_calls.values():
        unique_shape_count += len(
            {_unique_shape_key(call) for call in operator_calls}
        )
    input_tensor_count = sum(
        len(call.inputs)
        for call in collector.calls
    )
    _validate_row_limits(
        {
            "summary": 1 + len(grouped_calls),
            "unique_shapes": 1 + unique_shape_count,
            "detail": 1 + len(collector.calls),
            "input_tensors": 1 + input_tensor_count,
            "metadata": 1 + len(report_metadata) + 2,
        }
    )

    workbook = Workbook(write_only=True)
    stats = _ExportStats()
    _write_summary_sheet(workbook, collector.calls, stats)
    _write_unique_shapes_sheet(workbook, collector.calls, stats)
    _write_detail_sheet(workbook, collector.calls, iteration, rank, stats)
    _write_input_tensors_sheet(workbook, collector.calls, iteration, rank, stats)
    report_metadata["truncated_cells"] = stats.truncated_cells
    report_metadata["truncated_sheets"] = sorted(stats.truncated_sheets)
    _write_metadata_sheet(workbook, report_metadata, stats)
    workbook.save(destination)
    return destination
